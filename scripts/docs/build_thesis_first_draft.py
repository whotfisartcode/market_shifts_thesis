from __future__ import annotations

import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "thesis_academic_manuscript_no_title_page.docx"


BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(14)
TABLE_SIZE = Pt(10.5)
SMALL_SIZE = Pt(11)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, size=TABLE_SIZE) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = BODY_FONT
    run.font.size = size
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_text)
    run._r.append(fld_end)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Inches(0.49)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = BODY_FONT
        style.font.size = BODY_SIZE
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    for style_name in ["List Number", "List Bullet"]:
        style = styles[style_name]
        style.font.name = BODY_FONT
        style.font.size = BODY_SIZE
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(3)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.name = BODY_FONT
        run.font.size = BODY_SIZE
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_para(doc: Document, text: str, *, italic: bool = False, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Inches(0.49)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = BODY_FONT
    run.font.size = BODY_SIZE
    run.italic = italic
    run.bold = bold


def add_small_para(doc: Document, text: str, *, italic: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = BODY_FONT
    run.font.size = SMALL_SIZE
    run.italic = italic


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.5
    for run in p.runs:
        run.font.name = BODY_FONT
        run.font.size = BODY_SIZE
    p.add_run(text).font.size = BODY_SIZE


def add_numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = BODY_FONT
    run.font.size = BODY_SIZE


def add_table(doc: Document, headers: list[str], rows: list[list[str]], source: str | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, header, bold=True, size=TABLE_SIZE)
        set_cell_shading(cell, "EDEDED")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=TABLE_SIZE)
    if source:
        add_small_para(doc, f"Note: {source}", italic=True)


def add_toc_field(doc: Document) -> None:
    entries = [
        "Abstract",
        "List of Abbreviations",
        "1. Introduction",
        "1.1 Research Design Overview",
        "2. Literature Review",
        "3. Research Methodology",
        "3.1 Data Sources Used in the Empirical Artifact",
        "3.2 Current Production Panel",
        "3.3 Target Hierarchy and Label Counts",
        "4. Results",
        "4.1 Baseline Temporal Test Metrics",
        "4.2 P0 Audit Summary",
        "4.3 Economic Feature-Group Interpretation",
        "5. Discussion",
        "Conclusion",
        "References",
        "Appendices",
    ]
    for entry in entries:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(entry)
        run.font.name = BODY_FONT
        run.font.size = BODY_SIZE


def clean_markdown(text: str) -> str:
    text = text.replace("*", "")
    text = text.replace("`", "")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def load_references() -> list[str]:
    text = (ROOT / "docs" / "references.bib").read_text(encoding="utf-8")
    entries = re.split(r"\n(?=@)", text.strip())
    refs: list[str] = []
    for entry in entries:
        fields: dict[str, str] = {}
        kind_match = re.match(r"@(\w+)\{([^,]+),", entry.strip())
        if not kind_match:
            continue
        kind = kind_match.group(1).lower()
        for match in re.finditer(r"\n\s*(\w+)\s*=\s*\{(.*?)\},?", entry, re.DOTALL):
            value = re.sub(r"\s+", " ", match.group(2)).strip()
            fields[match.group(1).lower()] = value
        author = fields.get("author", "Unknown author").replace(" and ", "; ")
        year = fields.get("year", "n.d.")
        title = fields.get("title", "Untitled")
        venue = (
            fields.get("journal")
            or fields.get("booktitle")
            or fields.get("publisher")
            or fields.get("organization")
            or fields.get("institution")
            or fields.get("howpublished")
            or kind.capitalize()
        )
        details = []
        if fields.get("volume"):
            details.append(fields["volume"])
        if fields.get("number"):
            details.append(f"({fields['number']})")
        if fields.get("pages"):
            details.append(f"pp. {fields['pages'].replace('--', '-')}")
        suffix = ", ".join(details)
        doi_or_url = fields.get("doi") or fields.get("url")
        ref = f"{author} ({year}). {title}. {venue}"
        if suffix:
            ref += f", {suffix}"
        if doi_or_url:
            ref += f". {doi_or_url}"
        refs.append(ref)
    return refs


def read_csv_rows(rel_path: str) -> list[dict[str, str]]:
    with (ROOT / rel_path).open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


PRIMARY_TARGET_TABLE = [
    ["distress_next_4q", "31,702", "207", "0", "Strict legal distress benchmark"],
    ["failure_pressure_conservative_v2_next_4obs", "29,805", "3,003", "1,897", "Main broader failure-pressure target"],
    ["success_resilience_next_4q", "20,152", "12,519", "11,550", "Main success/resilience target"],
    ["industry_relative_resilience_next_4obs", "20,039", "8,093", "11,663", "Validated secondary production outcome"],
    ["stress_resilience_next_4obs", "5,901", "3,707", "25,801", "Validated secondary production outcome"],
    ["recovery_next_4obs", "11,375", "5,578", "20,327", "Validated secondary production outcome"],
    ["quality_success_cashflow_next_4obs", "16,767", "9,852", "14,935", "Validated secondary production outcome"],
]


PRIMARY_METRICS_TABLE = [
    ["distress_next_4q", "Gradient boosting", "5,872", "96", "0.104", "0.182"],
    ["failure_pressure_conservative_v2_next_4obs", "Random forest", "5,735", "575", "0.758", "0.725"],
    ["success_resilience_next_4q", "Random forest", "4,047", "2,648", "0.975", "0.938"],
    ["industry_relative_resilience_next_4obs", "Gradient boosting", "4,098", "1,678", "0.855", "0.787"],
    ["stress_resilience_next_4obs", "Random forest", "3,028", "1,920", "0.974", "0.937"],
    ["recovery_next_4obs", "Random forest", "2,419", "1,199", "0.977", "0.915"],
    ["quality_success_cashflow_next_4obs", "Gradient boosting", "3,998", "2,458", "0.953", "0.920"],
]


FEATURE_GROUP_TABLE = [
    ["distress_next_4q", "Gradient boosting", "Accounting fundamentals", "53.2%"],
    ["distress_next_4q", "Gradient boosting", "Firm trend / deterioration", "27.9%"],
    ["distress_next_4q", "Gradient boosting", "Financial ratios", "16.7%"],
    ["failure_pressure_conservative_v2_next_4obs", "Random forest", "Accounting fundamentals", "36.9%"],
    ["failure_pressure_conservative_v2_next_4obs", "Random forest", "Firm trend / deterioration", "34.6%"],
    ["failure_pressure_conservative_v2_next_4obs", "Random forest", "Financial ratios", "27.3%"],
    ["success_resilience_next_4q", "Random forest", "Financial ratios", "43.2%"],
    ["success_resilience_next_4q", "Random forest", "Firm trend / deterioration", "29.1%"],
    ["success_resilience_next_4q", "Random forest", "Accounting fundamentals", "26.2%"],
]


AUDIT_TABLE = [
    ["Prediction timestamp", "PASS", "All rows use filing-date prediction timestamps in the current panel."],
    ["Distress event dates", "PASS", "Formal strict-distress rows have source-provenance records."],
    ["Leakage", "PASS", "Target and metadata leakage fields are excluded from model features."],
    ["Macro lag", "PASS", "Macro variables are aligned to prediction dates."],
    ["Global-event timing", "PASS", "Event fields are attached through observed event windows."],
    ["Post-event rows", "PASS", "Post-event rows are retained for history but excluded from primary forward models."],
]


MAIN_BODY: list[tuple[str, int, list[str]]] = [
    (
        "1. Introduction",
        1,
        [
            "The problem addressed in this thesis is the early identification and interpretation of firm success and failure factors during changing market conditions. Public companies disclose large volumes of financial information, while macroeconomic agencies publish high-frequency series that describe the surrounding environment. The practical challenge is not a lack of data. The practical challenge is that the data must be converted into a temporally valid firm-period panel, linked to defensible future outcomes, and interpreted in a way that can support management, investor, and analytical decisions without turning statistical association into causal certainty.",
            "The thesis title, Navigating Market Shifts: Predictive Insights into Success and Failure Factors Across Industries, requires a broader view than a conventional bankruptcy-prediction exercise. Companies may experience formal legal distress, broader financial pressure, recovery, stable profitability, or sector-relative resilience. These outcomes are related, but they are not identical. A firm can become financially pressured without entering bankruptcy, and another firm can remain profitable in a favorable sector while still losing relative position against its peers. The thesis therefore treats success and failure as a hierarchy of observable forward-looking labels rather than as one overloaded binary outcome.",
            "The relevance of the topic follows from three connected developments. First, corporate financial distress remains important after the global financial crisis, the pandemic shock, the inflation and rate-hiking cycle, banking-stress episodes, supply-chain disruption, and energy-market volatility. Second, business analytics has made it easier to build large predictive systems, but not necessarily easier to make them auditable. Third, decision makers increasingly need explanations, not only scores. A model that ranks firms by future pressure is more useful when the user can see whether the signal comes from profitability, leverage, cash flow, deterioration, industry context, or macro-regime conditions.",
            "Prior literature gives strong foundations for this problem. Beaver (1966), Altman (1968), Ohlson (1980), Zmijewski (1984), Shumway (2001), and Campbell, Hilscher and Szilagyi (2008) show that accounting ratios, profitability, leverage, and dynamic risk measures can predict distress. More recent machine-learning literature, including Barboza, Kimura and Altman (2017), Dasilas and Rigani (2024), Zhao, Ouenniche and De Smedt (2024), and Kim, Cho and Ryu (2020), confirms that bankruptcy and financial-distress prediction remain active research areas. However, the maturity of the literature also means that a thesis cannot claim novelty simply by training another classifier on accounting ratios.",
            "The gap addressed here is narrower and more practical. Many studies focus on one distress definition, one sample design, or one table of model metrics. This thesis instead builds and evaluates an integrated business-analytics artifact: an audited SEC Financial Statement Data Sets and Federal Reserve Economic Data firm-period panel, a target hierarchy that separates strict legal distress from broader failure pressure and success/resilience, temporal validation based on filing-date prediction timestamps, leakage and data-quality audits, interpretable feature-group analysis, and a Streamlit dashboard for historical exploration and decision support.",
            "The research problem can be formulated as follows: how can a reproducible, temporally valid, and interpretable predictive-analytics framework be used to identify and compare observable success and failure factors across industries under market-shift conditions? The wording is intentional. The thesis does not seek to prove the complete causal theory of why companies fail. It seeks to construct a defensible empirical framework for forward-looking classification, factor interpretation, and artifact-based exploration.",
            "The goal of the thesis is to develop and evaluate an audited predictive-analytics framework for studying firm failure pressure, strict legal distress, and financial success/resilience across U.S. public companies and market environments. The analytical output should be useful both as an empirical study and as a decision-support artifact. The empirical contribution is the structured panel and target hierarchy; the practical contribution is a dashboard and set of interpretable outputs that make the model results inspectable by target, sector, firm, feature family, and market context.",
            "The research tasks are: to critically review literature on strategic adaptation, financial distress prediction, macro-financial market shifts, rare-event evaluation, interpretability, missing data, and design science; to construct a firm-period panel from SEC accounting data, FRED macro variables, industry metadata, global-event windows, and curated distress labels; to define strict distress, broader financial pressure, success/resilience, and validated secondary outcomes; to train and compare interpretable temporal models; to audit leakage, prediction timestamps, macro timing, event timing, selected SEC facts, and missingness; to interpret feature groups and target results; and to convert the empirical output into a dashboard-oriented decision-support narrative.",
            "The object of the research is the financial condition and forward outcomes of public companies observed through SEC filings across multiple industries and changing market environments. The subject of the research is the relationship between observable accounting fundamentals, derived financial ratios, firm deterioration features, macro-regime context, global-event windows, industry metadata, and forward-looking distress or resilience labels. This distinction matters because the study does not observe all managerial capabilities, private information, or strategic choices. It observes structured financial and contextual signals available from reproducible data sources.",
            "The main methods are literature analysis, data engineering, feature construction, temporal train-validation-test splitting, supervised machine learning, rare-event classification metrics, calibration and ranking diagnostics, leakage auditing, missingness diagnostics, feature-importance grouping, and dashboard-oriented design-science evaluation. The modeling design uses logistic regression, random forest, gradient boosting, and robustness benchmarks, while the interpretation emphasizes precision-recall area, F1, threshold behavior, and feature-group contribution rather than accuracy alone.",
            "The main information sources are SEC Financial Statement Data Sets, SEC EDGAR metadata, FRED macro-financial time series, a curated global-event calendar, a manually maintained and source-verified strict distress-event file, project audit outputs, model reports, and a curated literature base of 67 verified references. The production panel currently contains 31,702 rows, 190 columns, 540 SEC CIKs, and 540 ticker or display identifiers. Period dates run from 2009-03-31 to 2026-02-28, while prediction timestamps run from 2009-04-15 to 2026-03-31.",
            "The scientific novelty is not a new bankruptcy algorithm. The defensible novelty is the integration of a filing-date-aligned SEC/FRED/global-event panel, a strict-distress and broader-outcome target hierarchy, temporal validation, leakage and selected-fact provenance audits, missingness-aware target logic, interpretable feature-group analysis, and a dashboard artifact. The practical significance is that the output can support historical analysis of financial pressure and resilience, sector comparison, latest-filing early-warning views, and transparent discussion of limitations.",
            "The thesis is written in research format. Chapter 1 states the problem, relevance, goal, tasks, object, subject, methods, sources, novelty, practical significance, and structure. Chapter 2 reviews the theoretical and empirical literature. Chapter 3 describes the data architecture, feature design, target construction, modeling design, validation, and dashboard artifact. Chapter 4 presents empirical results. Chapter 5 discusses interpretation, managerial implications, limitations, and future research. The conclusion summarizes the answer to the research problem. Appendices contain supporting data dictionaries, target definitions, audit summaries, and dashboard notes.",
            "A central claim of the thesis is that target definition is an empirical design choice, not an administrative afterthought. Strict legal distress is clean and important, but it is rare. Broader financial pressure is not bankruptcy, but it is better suited for factor analysis because it captures repeated serious future weakness. Success/resilience is a positive counterpart that allows the thesis to study not only negative outcomes but also financially durable outcomes. Secondary validated targets add recovery, stress resilience, sector-relative resilience, and cash-flow-supported success as dimensional views.",
            "Another central claim is that temporal validity must be designed before modeling. In the production panel, prediction_date equals filed_date for every row. This prevents the model from acting as if a filing-period end date were the moment when the financial statement became known. It also creates a coherent basis for aligning macro variables, event windows, train-test splits, and future labels. This design is more conservative than a simple period-end panel and is more appropriate for an early-warning artifact.",
            "The thesis also takes a conservative view of missing values. Missing SEC-derived fields are not automatically zeros, not automatically unchanged prior values, and not automatically economic signals. A missing value may mean non-reporting, non-applicability, tagging variation, dimensional detail excluded from the standardized panel, or derivation limits. The modeling pipeline may impute within the training process, but the exported panel preserves null values. Missingness indicators may help prediction but are excluded from economic-factor interpretation.",
            "The expected contribution is therefore a disciplined empirical framework rather than an overclaim. The thesis can safely state that firm fundamentals, ratios, and deterioration features provide the main factor-interpretation layer in the current results, while macro regimes and global-event fields provide market-shift context and dashboard comparison dimensions. It cannot safely claim that feature importance proves causality, that broader financial pressure equals bankruptcy, or that the dashboard is a live daily bankruptcy oracle.",
        ],
    ),
    (
        "2. Literature Review",
        1,
        [
            "The literature review is organized around five layers: strategic adaptation and resilience, classical financial distress prediction, machine-learning prediction and evaluation, macro-financial market shifts, and interpretable decision-support artifacts. This structure follows the thesis problem. The research is not only a finance model, not only a machine-learning benchmark, and not only a dashboard project. It is an applied business-analytics study that combines theory, data engineering, predictive modeling, interpretation, and artifact design.",
            "Strategic-management theory motivates the broad language of success, failure, and market shifts. Barney (1991) argues that firm resources can explain persistent performance differences when resources are valuable, rare, difficult to imitate, and difficult to substitute. Teece, Pisano and Shuen (1997) extend this reasoning toward dynamic capabilities, which matter when environments change and firms must sense, seize, and reconfigure. Hannan and Freeman (1984) provide a counterweight by emphasizing structural inertia and selection pressures. Together, these theories justify studying why firms react differently under changing conditions, while also warning that the current panel can only observe financial proxies, not the full capability system inside each company.",
            "Resilience literature further supports the thesis framing, but it also requires caution. Linnenluecke (2017) shows that resilience is a broad construct across management research. Duchek (2020) frames organizational resilience as a capability-based process, and Hepfer and Lawrence (2022) emphasize that resilience can be functional, operational, or strategic. The present thesis does not measure total organizational resilience. It measures financial resilience through forward profitability, ROA, leverage, cash flow, recovery, and sector-relative performance proxies. This narrower interpretation makes the empirical claim defensible.",
            "Classical financial-distress literature remains foundational. Beaver (1966) demonstrated that financial ratios can carry early-warning information. Altman (1968) combined multiple financial ratios into the Z-score tradition. Ohlson (1980) introduced a probabilistic bankruptcy-prediction model based on accounting data. Zmijewski (1984) emphasized sample-selection and estimation issues. These studies justify the use of profitability, liquidity, leverage, activity, and solvency indicators. They also show that accounting-based distress prediction is not new, which means the thesis contribution must come from integration, target hierarchy, validation, auditing, and interpretability.",
            "Later distress-prediction research moved the field beyond static ratio snapshots. Shumway (2001) argued for dynamic hazard-style thinking because firms are observed repeatedly over time. Campbell, Hilscher and Szilagyi (2008) linked distress risk to profitability, leverage, market variables, and dynamic conditions. Tinoco and Wilson (2013) showed the value of accounting, market, and macroeconomic variables for listed-company distress and bankruptcy prediction. Duffie, Saita and Wang (2007) also support the use of multi-period covariates in default prediction. These sources support the present firm-period design and the use of macro context, while leaving room for future market-price extensions not implemented in the current panel.",
            "Market-based default literature is important but mostly outside the implemented scope. Merton (1974) created a structural model connecting corporate debt risk to firm asset value and volatility. Hillegeist et al. (2004) and Bharath and Shumway (2008) discuss market-based bankruptcy probability and distance-to-default style measures. These sources are relevant to limitations and future work because they show an alternative data tradition. The current production panel is intentionally SEC/FRED based and does not implement a full market-value default model, so those papers should not be cited as if their method were used directly.",
            "Modern machine-learning literature confirms that nonlinear models can be useful for financial distress and related risk tasks. Lessmann et al. (2015) benchmarked classification algorithms for credit scoring. Barboza, Kimura and Altman (2017) found that machine-learning models can improve bankruptcy prediction relative to some traditional methods. Dasilas and Rigani (2024), Zhao, Ouenniche and De Smedt (2024), and Kim, Cho and Ryu (2020) show that ML-based distress, bankruptcy, and corporate default prediction remain current topics. These sources support model comparison but also reinforce that algorithmic performance depends on target definition, validation design, and data quality.",
            "The implemented model families are standard rather than exotic. Breiman (2001) provides the foundation for random forests, Friedman (2001) for gradient boosting, and Pedregosa et al. (2011) for the scikit-learn implementation ecosystem. Their inclusion supports reproducibility and method description. The thesis does not need to claim that a novel classifier has been invented. It needs to show that known model families are used in a controlled, auditable, temporally valid way on a carefully constructed panel.",
            "Rare-event evaluation is central because strict legal distress is sparse. Fawcett (2006) explains receiver operating characteristic analysis, but Davis and Goadrich (2006) and Saito and Rehmsmeier (2015) show why precision-recall curves are especially informative under class imbalance. A model can have a respectable ROC-AUC while still producing many false positives when the positive class is rare. This is why the thesis reports PR-AUC, precision, recall, F1, confusion matrices, test positives, and ranking behavior. Accuracy alone would be misleading.",
            "Class imbalance literature also matters. Chawla et al. (2002) introduced SMOTE, and Burez and Van den Poel (2009) discuss class-imbalance handling in predictive modeling. The thesis does not make synthetic oversampling the main story because time-ordered financial panels can be distorted by careless resampling. Instead, it emphasizes target hierarchy, rare-event metrics, temporal validation, and threshold diagnostics. This position is conservative and easier to defend.",
            "Macro-financial literature explains why market shifts belong in the thesis. Bernanke, Gertler and Gilchrist (1999) describe the financial accelerator, where credit frictions amplify macroeconomic shocks. Gilchrist and Zakrajsek (2012) show that credit spreads contain important business-cycle information. Longstaff, Mithal and Neis (2005) caution that spreads combine default and liquidity components, which means interpretation must be careful. These sources support the inclusion of rates, spreads, volatility, financial conditions, inflation, oil shocks, demand slowdown, credit tightening, and crisis indicators as context variables.",
            "Official data documentation also belongs in the literature base because the thesis is partly an artifact-building study. SEC Financial Statement Data Sets and EDGAR API documentation support the accounting and filing source. FRED documentation supports macro-series retrieval and interpretation. Federal Reserve financial-stability reports and Chicago Fed NFCI documentation support market-condition context. These sources are not theory papers, but they are necessary for reproducibility and data lineage.",
            "The literature on missing data supports the project's conservative null-treatment policy. Rubin (1976), Little and Rubin (2019), Graham (2009), and van Buuren (2018) show that missingness mechanisms affect inference and should not be ignored. In SEC data, a missing value may be structural, reporting-related, tag-related, derivation-related, or genuinely unavailable. The thesis therefore separates raw null-preserving data from model-estimation imputation. This prevents missingness indicators from being described as economic causes of failure or success.",
            "Interpretability literature is needed because the thesis promises predictive insights, not merely predictive scores. Lundberg and Lee (2017) introduced SHAP as a unified explanation framework, Ribeiro, Singh and Guestrin (2016) introduced local explanation logic, and Guidotti et al. (2018) survey black-box explanation methods. Finance-specific work by Bussmann et al. (2021), Gramegna and Giudici (2021), and Zhang et al. (2022) supports the importance of explainable models in financial risk settings. Molnar (2025) provides practical caution: explanations describe model behavior and should not be treated as proof of causality.",
            "Design-science and business-analytics literature supports the dashboard artifact. Hevner et al. (2004) and Peffers et al. (2007) justify the construction and evaluation of information-system artifacts. Chen, Chiang and Storey (2012) position business intelligence and analytics as a field concerned with data, models, and decision impact. Shneiderman (1996) supports interactive visual exploration principles, and Arnott and Pervan (2014) connect decision support systems with design-science logic. These sources help frame the dashboard as part of the thesis contribution rather than as a decorative add-on.",
            "Emerging-market and Russia-specific sources are useful for boundary conditions. Charalambakis and Garrett (2015), Hunter and Isachenkova (2001), Afanasev (2023), Fedorova, Musienko and Fedorov (2020), and Nevredinov (2021) indicate that distress predictors and data availability can differ across institutional settings. The present empirical scope remains U.S. public companies because SEC and FRED data are reproducible and audited in the project. Russian or broader international data are better treated as future research unless a separate validated panel is built.",
            "The literature synthesis leads to a precise gap. Existing research has strong classical distress models, modern machine-learning benchmarks, macro-default studies, explainability methods, and dashboard methodology. What remains useful is the integration of these elements into one auditable business-analytics artifact with filing-date timestamps, multiple forward-looking target definitions, temporal validation, leakage controls, missingness policy, factor-group interpretation, and dashboard exploration. This is the gap that the thesis addresses.",
            "The author's position is that a defensible predictive study of success and failure must distinguish three levels: label validity, model validity, and interpretation validity. Label validity asks whether the future outcome means what the text says it means. Model validity asks whether the validation design reflects information available at prediction time. Interpretation validity asks whether feature importance is explained as association within the model rather than causal proof. The current project is organized around these three levels.",
            "The literature also supports the decision not to add uncontrolled scope. A full distance-to-default model, large-scale NLP pipeline, cross-country accounting harmonization, or live global event feed could be valuable, but each would require separate data lineage, validation, and limitations. The thesis instead focuses on a cleaner implemented core. That choice is not a weakness; it is a methodological boundary that allows the study to make stronger claims about the data and results it actually contains.",
            "From the literature, four research propositions guide the empirical part. First, firm fundamentals, ratios, and deterioration features should carry substantial predictive information for financial pressure and resilience. Second, strict legal distress should be harder to predict than broader pressure because it is rare and legally specific. Third, macro-regime and global-event variables should add context and support sector comparisons even if they do not dominate firm accounting signals. Fourth, interpretable feature groups and dashboards should improve the practical usability of predictive outputs.",
        ],
    ),
    (
        "3. Research Methodology",
        1,
        [
            "The research design is an applied predictive-analytics and design-science study. It creates a reproducible empirical artifact, evaluates it through statistical models and audits, and presents the output through a dashboard. The unit of observation is one firm-period row: one public company, one reported accounting period, observed from the date on which the SEC filing became available. This unit is appropriate because the thesis studies forward outcomes after information becomes public, not only accounting states at fiscal period end.",
            "The production panel combines four source families. The first is SEC Financial Statement Data Sets, which provide structured accounting facts from public-company filings. The second is FRED, which provides macroeconomic and financial-market variables. The third is a curated global-event calendar, which maps historical market-shift windows to dates and event types. The fourth is a manually curated distress-event file with source provenance for formal strict-distress rows. Firm and industry metadata are attached through SEC identifiers, ticker/display fields, SIC information, and sector classifications.",
            "The current production panel has 31,702 rows, 190 columns, 540 SEC CIKs, and 540 ticker or display identifiers. The period range is 2009-03-31 to 2026-02-28, and the prediction timestamp range is 2009-04-15 to 2026-03-31. The prediction-date policy is the central integrity decision: prediction_date equals filed_date for every row in the current panel. This means the model does not treat period-end statements as known before filing. The policy also aligns macro variables, event windows, temporal splits, and future target labels to the same information date.",
            "SEC accounting facts are processed conservatively. Balance-sheet variables use point-in-time values. Flow variables use single-period values when directly reported. If a firm reports cumulative year-to-date flow values, the builder derives the single-period value by subtracting prior cumulative values within the same firm, fiscal year, and variable. If the required prior cumulative fact is unavailable, the value remains null. This treatment improves coverage without inventing data and preserves the distinction between reported and derived accounting values.",
            "The feature set is organized into interpretable families. Raw accounting fundamentals include total assets, total liabilities, current assets, cash equivalents, revenue, operating income, net income, operating cash flow, and capital expenditure. Financial ratios include ROA, leverage/assets, current ratio, cash/assets, net margin, operating margin, gross margin, R&D intensity, inventory/assets, and receivables/assets. Firm trend and deterioration features include lags, four-observation changes, growth rates, and prior negative-income counts. These features are derived from current or prior observations, not future data.",
            "Macro variables come from FRED and include interest rates, Treasury spreads, VIX, credit spreads, financial conditions, money supply, industrial production, inflation, oil prices, dollar strength, employment, retail sales, housing starts, business lending, and lending-standard tightening. Year-over-year changes are derived for selected level and index series. Regime indicators convert continuous macro data into interpretable market states such as high-rate periods, market-stress periods, tight financial conditions, credit-spread stress, yield-curve inversion, inflation pressure, oil shocks, strong-dollar regimes, demand slowdown, credit tightening, and crisis years.",
            "Global-event context is added through a curated calendar rather than a raw news pipeline. Event types include financial crisis, sovereign debt, natural disaster, commodity oil, financial market, political policy, trade policy, pandemic, supply chain, monetary inflation, geopolitical war, and banking stress. The constructed event columns include event counts, event severity sums, event-type counts, and event names. This layer supports market-shift interpretation and dashboard filtering, but it is not described as a complete event database or live news model.",
            "The universe combines the defended-thesis core, additional controls, distressed and near-distressed candidates, event-review expansion firms, and matched controls. The final production panel contains 540 SEC CIKs after coverage, identifier, alias-collapse, and panel-construction constraints. The old 255-ticker legacy panel is treated only as history because it had data-quality and methodology caveats. The current panel is the empirical basis of the thesis.",
            "Target construction follows a hierarchy. The strict legal distress benchmark, distress_next_4q, is positive when a formal distress event occurs after prediction_date and within 456 days. It is clean but rare. The main failure-factor target, failure_pressure_conservative_v2_next_4obs, is a broader conservative pressure label. It captures strict formal distress or repeated serious future profitability weakness combined with balance stress, deterioration, or unhealthy future signals. It is not a bankruptcy label. The main success target, success_resilience_next_4q, is positive when future observations show positive net income, positive ROA, acceptable leverage/assets, and no formal distress.",
            "The target hierarchy also includes four validated secondary production outcomes: industry_relative_resilience_next_4obs, stress_resilience_next_4obs, recovery_next_4obs, and quality_success_cashflow_next_4obs. These outcomes are not replacements for the three primary targets. They add dimensional analysis of sector-relative performance, resilience during stress conditions, recovery from weaker states, and cash-flow-supported success. Two additional robustness-extension targets, sector_relative_improvement_next_4obs and persistent_resilience_next_6obs, are discussed as sensitivity views rather than as production targets.",
            "Missing-aware target logic is important. Post-event rows are left null for broader forward-looking labels. Future outcome components remain unknown when too few future observations are available. Missing future liabilities or leverage values are not treated as automatic non-success. This is a correction from less conservative target logic and is one reason the current success/resilience denominator is smaller than the full panel. The denominator difference is a validity feature, not a defect.",
            "The modeling design uses temporal splits instead of random row splits. Training, validation, and test periods are defined by prediction dates so that later periods evaluate whether the model generalizes forward in time. Model rows exclude post-event observations for forward-looking prediction. The baseline model families are logistic regression with L2 regularization, random forest, and gradient boosting. Controlled tuning and advanced boosted-tree benchmarks are available as robustness evidence, but the final thesis story emphasizes transparent baseline comparison and temporal validity.",
            "The evaluation metrics match the target properties. ROC-AUC is reported because it is familiar, but PR-AUC, precision, recall, F1, test positives, threshold behavior, and top-rank precision are more informative for rare distress. Strict legal distress is evaluated as a rare-event benchmark. Broader failure pressure and success/resilience are evaluated as richer factor-analysis outcomes with larger positive classes. Metrics are not compared mechanically across targets because target prevalence and meaning differ.",
            "The model pipeline handles missing values inside estimation while preserving raw panel nulls. Numeric variables are imputed with training-split medians, and missingness indicators may be created inside the scikit-learn preprocessing pipeline. However, missingness indicators are excluded from economic feature-group interpretation. This separation allows the model to use reporting-pattern information when useful while preventing the thesis from claiming that missing fields are direct economic causes of success or failure.",
            "Leakage control uses explicit feature blacklists and audits. Target columns, future-looking fields, post-event indicators, event dates, days-to-event, sampling metadata, and other non-predictive identifiers are excluded from model features where they would leak labels or sampling logic. P0 audits check prediction timestamp validity, distress-event dates, leakage, macro lags, global-event timing, and post-event row treatment. The model-output consistency audit checks that saved model reports align with the production target set.",
            "Selected SEC fact provenance is also audited. The project verifies that selected SEC facts align to the intended period, with zero selected rows where ddate differs from period in the current provenance output and zero selected-value mismatches above tolerance. Same-information-date amendment duplicates are resolved deterministically before ratios, trends, and targets are built. A duplicate CIK issue involving PCG and PGNPQ was resolved through alias collapse, keeping one CIK history and preserving event-source alias metadata.",
            "The methodology includes a dashboard artifact because the thesis is positioned as business analytics and decision support. The dashboard exposes dataset composition, target coverage, model metrics, feature groups, firm exploration, macro comparison, event context, and artifact caveats. It is not a live trading system or daily bankruptcy oracle. It is a historical market-shift and latest-filing decision-support tool. Firm accounting signals update when new filings become available; macro and event context may update more frequently.",
            "The design-science evaluation of the dashboard is pragmatic. The dashboard is assessed by whether it separates targets from firm metrics, exposes data coverage and caveats, allows sector and regime comparison, displays model metrics and top feature groups, supports firm-level trajectory review, and preserves reproducibility notes. The screenshot capture report marks dashboard screenshot views as passing, with only known benign Vega-Lite warnings recorded. This supports artifact readiness without overstating scientific proof.",
            "The methodology is deliberately conservative in scope. It does not add RFSD, Hong Kong, Japan, paid default databases, large-scale NLP, live event feeds, GANs, LSTMs, transformers, HR data, ESG ratings, patents, or customer sentiment to the final empirical model. These ideas are relevant future work, but adding them before submission would create data-lineage and validation risks. The thesis focuses on implemented, audited, reproducible data and models.",
            "The practical implication of the methodology is that the empirical chapter can answer several linked questions. Which financial factors dominate strict distress, broader pressure, and success/resilience? How much harder is strict legal distress than broader pressure? Do secondary targets reveal recovery and stress-resilience dimensions? Do macro regimes and event windows contextualize sector behavior? Can the resulting information be explored through a dashboard without hiding target definitions or caveats?",
            "The limitation of this methodology is that it studies observable financial and contextual signals, not the full set of strategic causes. It does not observe management quality, private negotiations, contract structure, supply-chain dependencies, or unreported operational stress. It also does not prove causality. The strength of the design is that these limits are explicit and that every final empirical claim is tied to a source, target, validation design, audit, or saved model output.",
        ],
    ),
    (
        "4. Results",
        1,
        [
            "The empirical results begin with the data-construction outcome. The production panel contains 31,702 firm-period rows, 190 columns, 540 SEC CIKs, and 540 ticker or display identifiers. The period-date range is 2009-03-31 to 2026-02-28, and the prediction-date range is 2009-04-15 to 2026-03-31. Processed and GitHub-ready panel copies match exactly. The panel therefore provides a single rectangular dataset linking SEC accounting fundamentals, derived ratios, trend features, macro variables, regime indicators, global-event context, industry metadata, and target labels.",
            "The audit results support structural readiness for thesis writing. P0 audits pass for prediction timestamp, distress-event dates, leakage, macro lag, global-event timing, and post-event rows. Selected SEC fact provenance contains 700,844 selected facts, zero selected facts with ddate not equal to period, and zero selected-value mismatches above tolerance. Deterministic duplicate handling dropped 15 same-information-date amended rows, leaving zero duplicate ticker/period/prediction groups. Numeric infinities are zero. These results do not prove causal truth, but they support data integrity and reproducibility.",
            "The target composition confirms why a hierarchy is necessary. Strict legal distress has 31,702 known rows and 207 positives. It is clean but rare. Broader failure pressure has 29,805 known rows, 3,003 positives, and 1,897 missing rows. Success/resilience has 20,152 known rows, 12,519 positives, and 11,550 missing rows. The missing rows are not errors; they reflect missing-aware future horizon rules. The validated secondary outcomes have different denominators because they answer different questions about relative resilience, stress resilience, recovery, and cash-flow-supported success.",
            "The strict legal distress benchmark remains difficult. The accepted baseline model is gradient boosting, with 5,872 test rows, 96 test positives, test PR-AUC of 0.104, and F1 of 0.182. These values should not be dismissed as failure without context. The positive class is only a small share of the test period, and the label is formal legal distress within a forward horizon. The model is useful as a clean benchmark and early-warning proof of concept, not as the headline performance story.",
            "The broader failure-pressure target produces a stronger factor-analysis model. The baseline random forest has 5,735 test rows, 575 test positives, test PR-AUC of 0.758, and F1 of 0.725. This target is more suitable for identifying failure factors because it captures serious repeated financial pressure before or outside formal legal distress. Its stronger metrics do not mean it predicts bankruptcy; they mean broader deterioration is more frequent and more directly reflected in accounting and deterioration features.",
            "The success/resilience target is empirically strong. The baseline random forest has 4,047 test rows, 2,648 test positives, test PR-AUC of 0.975, and F1 of 0.938. This result should be interpreted carefully because the target is not rare and because positive success/resilience labels require observed future profitability, ROA, leverage/assets, and no formal distress. The high performance indicates that observable financial condition and recent trends are strongly associated with future financially resilient states under the target definition.",
            "The four validated secondary production outcomes add useful dimensions. Industry-relative resilience, stress resilience, recovery, and quality-success-cashflow targets all passed profile, temporal model sanity, calibration, ranking, and interpretation checks before promotion. Baseline test PR-AUC values are 0.855, 0.974, 0.977, and 0.953 respectively, with strong F1 values. These outcomes should be presented as secondary analysis, not as replacements for the three primary targets. They help answer the broader thesis title by showing different success and recovery concepts.",
            "Feature-group interpretation supports the central substantive result. For strict distress under gradient boosting, accounting fundamentals account for about 53.2% of economic importance after excluding missingness indicators, firm trend/deterioration for about 27.9%, and financial ratios for about 16.7%. For broader failure pressure under random forest, accounting fundamentals, firm trend/deterioration, and financial ratios account for about 36.9%, 34.6%, and 27.3% respectively. For success/resilience under random forest, financial ratios, trend/deterioration, and accounting fundamentals account for about 43.2%, 29.1%, and 26.2% respectively.",
            "The feature-group result is more useful than isolated feature ranking because it is closer to economic interpretation. It shows that the main signals are not random identifiers or hidden target leakage. They are accounting facts, ratios, and deterioration patterns. Industry and macro context appear as supporting layers rather than dominant drivers in the main tree-based interpretation. This finding is consistent with classical distress literature and with the project's conservative claim guardrails.",
            "Ablation evidence also supports this conclusion. In the strict-distress ablation output, random forest with firm plus macro/regime features reaches higher PR-AUC than firm-only in one specification, while firm-only has strong ROC-AUC. Event-aware features are useful for dashboard and context but do not replace firm fundamentals. The correct interpretation is that market-shift variables help frame and segment the analysis, while firm-level financial condition carries the primary signal for the current targets.",
            "Calibration and ranking outputs add practical meaning. For strict distress, top-ranked groups have much higher precision than the base rate, but recall remains constrained because the class is rare. For broader failure pressure, top 10% ranking in the test split captures a large share of positive cases with high precision relative to the base rate. For success/resilience, higher probability thresholds identify observations with high precision, but the base rate is already high. These threshold results support dashboard ranking and alert use, not automatic decision-making.",
            "The results also show why target denominators must be disclosed. Success and secondary targets often have many missing future labels because the required future observations are not always available. Treating those missing futures as negatives would artificially change target meaning. The current missing-aware policy is stricter and more scientifically defensible. It makes the dataset less superficially complete but better aligned with the outcome definitions.",
            "Sector and event summaries support the 'across industries' and 'market shifts' language. Strict distress positives concentrate more visibly in sectors such as Consumer Cyclical and Energy in the available summaries, while other sectors show fewer strict events. However, sector interpretation must be cautious because sector sizes and event counts differ. The dashboard is useful here because it allows the reader to compare denominators, positive shares, and firm examples rather than relying on a single aggregate statement.",
            "The dashboard artifact is a result in its own right. It presents dataset and target composition, coverage and missingness, primary model metrics, target-lab outcomes, feature groups, firm explorer views, macro comparisons, target timelines, and artifact caveats. The firm explorer separates financial metrics from target timelines so scales are not confused. The macro comparison uses a training-period baseline. The caveat panel keeps scope, data lineage, and claim boundaries visible to the reader.",
            "The final empirical story is therefore balanced. Strict legal distress is clean but rare and remains a difficult benchmark. Broader failure pressure is the main failure-factor outcome because it has more positive cases and captures serious financial deterioration. Success/resilience is strong and interpretable under missing-aware future logic. Secondary outcomes add useful dimensions of relative resilience, stress resilience, recovery, and cash-flow-supported success. Across targets, firm fundamentals, ratios, and deterioration dominate economic interpretation, while macro regimes and global events contextualize market shifts.",
        ],
    ),
    (
        "5. Discussion",
        1,
        [
            "The first implication is that the thesis should not be presented as a simple bankruptcy-prediction success story. The strict legal distress target is important precisely because it is clean, sourced, and difficult. Its rarity means that even a model with useful ranking ability can have modest F1 and PR-AUC. This does not invalidate the project; it clarifies the role of the target. Strict distress functions as a benchmark and integrity anchor, while broader pressure and resilience targets carry the main factor-analysis story.",
            "The second implication is that target hierarchy improves scientific honesty. A single success/failure target would mix legal events, financial deterioration, recovery, and ordinary profitability into one ambiguous label. The hierarchy separates these concepts. It allows the thesis to say, for example, that broad financial pressure is detectable with stronger performance than strict legal distress, while success/resilience is a positive counterpart with its own denominator and meaning. This distinction is one of the core methodological contributions.",
            "The third implication is that firm-level accounting and deterioration remain central. The feature-group outputs show that fundamentals, ratios, and trends dominate the economic interpretation layer for the main tree-based models. This is consistent with the classical literature and with practical expectations: firms under pressure tend to show weakness in assets, liabilities, profitability, cash flow, revenue, margins, leverage, and deterioration patterns before or during stress. The model results update this tradition with an audited panel and temporal validation rather than replacing it with a purely macro or black-box explanation.",
            "The fourth implication is that macro regimes and global events should be framed as context, not as a replacement for firm analysis. Rates, inflation, credit spreads, volatility, oil shocks, crisis periods, supply-chain events, and banking stress all matter for interpretation. They make the title's 'market shifts' concrete and allow sector comparison. However, the current empirical results do not support a claim that macro variables alone dominate firm fundamentals. The dashboard should therefore use macro/event layers for explanation, segmentation, and historical exploration.",
            "The managerial value of the artifact is practical but bounded. A manager, analyst, investor, or supervisor can use the dashboard to inspect which firms show elevated pressure under the latest available filing, which feature groups drive model outputs, how sector outcomes differ across regimes, and how target definitions change denominators. The artifact can support triage, discussion, and monitoring. It should not be used as an automatic credit decision engine or daily legal-bankruptcy forecast.",
            "For organizations, the factor interpretation suggests that basic financial discipline remains visible in predictive systems. Profitability, leverage, cash availability, balance-sheet structure, operating performance, and deterioration trends appear repeatedly across targets. This does not mean managers can mechanically optimize a model score. It means that transparent monitoring of these areas can help identify pressure earlier and explain why a firm is classified as vulnerable, resilient, recovering, or sector-relative strong under the empirical definitions.",
            "For investors and market analysts, the target hierarchy creates a more nuanced screening language. Strict distress flags are rare and may be late or legally specific. Broader pressure identifies repeated serious weakness. Success/resilience identifies financially durable states. Secondary outcomes distinguish firms that improve relative to sector peers, recover after weakness, withstand stress contexts, or combine profitability with cash-flow support. This structure is more useful than a single binary label because it aligns the screening output with different analytical questions.",
            "For regulators and researchers, the thesis demonstrates the value of filing-date alignment and auditability. Many predictive systems fail silently because they leak future information through period dates, target-derived variables, post-event rows, or unexamined data joins. The current project explicitly audits prediction timestamps, macro lags, event timing, leakage, post-event rows, selected SEC fact provenance, duplicate handling, and package equality. These controls make the empirical results more defensible even when some target performance remains modest.",
            "The dashboard also has educational value. It makes the difference between data, target, model, and interpretation visible. Users can see that raw accounting fields differ from ratios, that model scores differ from realized target labels, that missingness affects target denominators, and that dashboard exploration is not the same as causal explanation. This transparency supports the design-science argument that the artifact improves understanding, not merely model deployment.",
            "Several limitations remain. First, the panel is U.S. public-company oriented. SEC data are reproducible and rich, but the findings cannot be mechanically generalized to private firms, Russian firms, emerging markets, or jurisdictions with different accounting and disclosure systems. Second, strict distress labels are source-verified for the current formal rows but still represent a curated event set, not an exhaustive legal database of every possible distress event. Third, broader failure pressure is an operational target, not legal bankruptcy.",
            "Fourth, feature importance is not causality. A feature group can be important because it helps the model separate future labels, but that does not prove that the feature caused the outcome. Profitability, leverage, and cash flow are economically meaningful, yet the model remains observational. Fifth, macro and event context may reflect broad conditions but cannot identify every channel through which market shifts affect firms. Sixth, missingness indicators may improve prediction but are not economic mechanisms.",
            "Seventh, the model set is intentionally conservative. Random forest and gradient boosting are powerful enough for the thesis, but the project does not implement all possible methods. Hazard models, distance-to-default, deep learning, graph models, text-based event extraction, and causal inference designs are future research paths. Eighth, the dashboard is an artifact based on latest available filings and saved project outputs; it is not a live production system with continuous data ingestion, governance, monitoring, and user testing.",
            "Future research should extend the panel with market-price variables and structural default measures, provided that timestamp alignment and coverage can be maintained. Distance-to-default and equity-volatility features could complement SEC accounting fundamentals, especially for listed firms. Another future path is a separate international or Russian panel with its own accounting harmonization, bankruptcy-label policy, and macro-data timing controls. Such work should not be merged casually into the current U.S. panel because institutional differences are material.",
            "A second future path is richer event data. GDELT, ACLED, EM-DAT, or other event sources could support automated event aggregation, but the data would need cleaning, category mapping, bias assessment, and date alignment. The current curated event calendar is intentionally smaller and more interpretable. A future live dashboard could update macro and event context more frequently while keeping firm fundamentals tied to filing dates.",
            "A third future path is more formal interpretability and robustness analysis. SHAP, permutation importance, partial dependence, model stability tests, and reason-code evaluation can deepen the explanation layer. The current thesis uses feature-group interpretation and saved importance outputs as a conservative layer. More advanced explanation tools would be useful, but they should retain the same caveat: explanations describe model behavior and should not be treated as causal proof.",
            "A fourth future path is decision-process evaluation. The dashboard could be tested with analysts, managers, or students to assess whether it improves understanding, reduces confusion about target definitions, or supports better triage decisions. That would move the artifact evaluation beyond structural readiness and screenshot checks toward user-centered decision-support evidence. Such testing is outside the current empirical scope but fits the business analytics program profile.",
            "Overall, the discussion supports a disciplined conclusion. The project is strongest when it is framed as an audited predictive-analytics framework for factor discovery and decision support. It is weaker if presented as a universal bankruptcy oracle. The current evidence supports the former framing: a reproducible panel, valid timestamps, target hierarchy, temporal models, audits, feature-group interpretation, and dashboard artifact that together provide predictive insights into success and failure factors across industries under market shifts.",
        ],
    ),
    (
        "Conclusion",
        1,
        [
            "The thesis set out to answer how a reproducible, temporally valid, and interpretable predictive-analytics framework can identify and compare observable success and failure factors across industries under market-shift conditions. The answer is that the framework must begin with auditable data and target design, not with model choice. The project builds a single firm-period panel from SEC accounting data, FRED macro variables, curated global-event context, industry metadata, and source-verified distress-event labels. It aligns prediction timestamps to SEC filing dates and separates strict distress, broader pressure, success/resilience, and secondary outcomes.",
            "The first task, literature review, shows that the topic rests on mature but fragmented foundations. Classical distress models justify financial ratios and accounting fundamentals. Dynamic and hazard-style research supports firm-period and temporal thinking. Macro-financial literature supports regime and market-shift context. Machine-learning and rare-event evaluation literature supports model comparison and PR-AUC. Missing-data literature supports conservative null treatment. Interpretability and design-science literature support feature-group explanation and dashboard design. The thesis contribution is not to replace these literatures, but to integrate them in an audited artifact.",
            "The second task, data construction, is fulfilled by the current production panel of 31,702 rows, 190 columns, 540 SEC CIKs, and 540 ticker/display identifiers. The panel preserves nulls, uses filing dates as prediction dates, improves SEC concept mapping conservatively, resolves amendment duplicates, resolves the PCG/PGNPQ alias issue, aligns selected SEC facts to periods, and keeps processed and GitHub-ready copies equal. P0 audits pass for prediction timestamp, distress-event dates, leakage, macro lag, global-event timing, and post-event rows.",
            "The third task, target construction, is fulfilled through the target hierarchy. Strict legal distress remains a clean rare-event benchmark with 207 positives. Broader conservative failure pressure becomes the main failure-factor target with 3,003 positives among 29,805 known rows. Success/resilience becomes the main positive outcome with 12,519 positives among 20,152 known rows. Four validated secondary production outcomes add industry-relative resilience, stress resilience, recovery, and cash-flow-supported success. This hierarchy makes the empirical story more honest and more useful.",
            "The fourth task, modeling and evaluation, shows that target meaning affects model performance. Strict legal distress is difficult, with gradient boosting test PR-AUC of 0.104 and F1 of 0.182 in the current baseline table. Broader failure pressure is stronger, with random forest test PR-AUC of 0.758 and F1 of 0.725. Success/resilience is stronger still, with random forest test PR-AUC of 0.975 and F1 of 0.938. Secondary targets also show strong temporal model performance. These results support factor analysis and dashboard ranking, not causal proof.",
            "The fifth task, interpretation, shows that accounting fundamentals, financial ratios, and firm deterioration features dominate the economic feature-group layer for the primary tree-based models. Macro regimes and global-event fields provide important market-shift context and dashboard segmentation, but they do not replace firm-level financial analysis. Missingness indicators are retained only as diagnosed predictive auxiliaries and are excluded from economic factor interpretation. This supports a clear and defensible practical message.",
            "The sixth task, artifact development, is fulfilled by the Streamlit dashboard and related screenshots/audits. The dashboard presents dataset composition, target coverage, model metrics, factor groups, firm explorer views, macro comparison, target timelines, and caveat notes. It is positioned as a historical market-shift and latest-filing decision-support artifact. It supports practical significance by making model outputs inspectable and by helping users understand the difference between data, target labels, predictions, and interpretation.",
            "The main scientific novelty is the integrated framework: an audited SEC/FRED/global-event firm-period panel, filing-date prediction timestamps, target hierarchy, temporal validation, leakage controls, selected-fact provenance, missingness-aware targets, interpretable factor-group analysis, and dashboard artifact. The main practical significance is a reproducible decision-support workflow for comparing failure pressure and success/resilience factors across firms, sectors, regimes, and event windows.",
            "The conclusion is therefore deliberately bounded. The thesis does not prove why every firm succeeds or fails, does not claim that broader pressure equals bankruptcy, does not claim that macro variables dominate firm fundamentals, and does not claim that feature importance is causal. It does show that a disciplined, auditable predictive-analytics workflow can produce useful and interpretable insights into success and failure factors across industries during market shifts.",
        ],
    ),
]


ACADEMIC_EXTENSIONS: dict[str, list[str]] = {
    "1. Introduction": [
        "The research problem is especially relevant because market shifts are no longer exceptional episodes that can be separated cleanly from ordinary business analysis. The sample period includes the aftermath of the global financial crisis, the long low-rate period, the COVID-19 shock, supply-chain disruption, inflation pressure, rapid monetary tightening, energy-price volatility, and banking-stress windows. These events affected firms unevenly. Some companies remained profitable and liquid, some recovered after temporary pressure, and some moved toward formal or economic distress. A useful empirical framework must therefore compare firms within a common structure while still allowing target definitions to differ by outcome.",
        "The topic also has practical importance for business analytics because decision makers often face a gap between descriptive financial reporting and forward-looking interpretation. Financial statements are historical documents, but they contain signals about future pressure or resilience. Macroeconomic indicators describe the environment, but they do not automatically identify which firms are vulnerable. Machine-learning models can rank observations, but they can also obscure target construction and data leakage. The thesis responds to this gap by designing a workflow in which data lineage, target logic, validation, and interpretation are all treated as parts of the empirical contribution.",
        "The degree of academic development is substantial but uneven across the elements combined in this thesis. Bankruptcy prediction has a long literature, macro-financial default modeling has a separate tradition, explainable artificial intelligence has developed rapidly, and design-science research supports dashboard artifacts. However, these streams are often implemented separately. A study may have strong classification metrics but weak timestamp controls; another may have a rich theoretical discussion but no reproducible data artifact; a third may include macro variables but use a single legal distress label. The present thesis occupies the intersection of these streams.",
        "The goal is formulated as an applied research goal rather than a purely technical model-development goal. The objective is to create and evaluate a reproducible, temporally valid, and interpretable framework for identifying observable financial success and failure factors across industries under market-shift conditions. This goal implies that model performance is necessary but not sufficient. The framework must also show how variables are constructed, how labels are defined, how future information is excluded, how missing values are treated, and how outputs can be understood by an analyst.",
        "The research tasks follow the logic of the artifact. First, the literature on strategic adaptation, financial distress prediction, macro-financial shocks, rare-event classification, interpretability, and decision support must be critically synthesized. Second, an empirical panel must be constructed from standardized accounting filings, macro variables, event windows, and industry metadata. Third, the outcomes must be defined as a hierarchy rather than a single success/failure label. Fourth, temporal models must be trained and evaluated with metrics appropriate to imbalanced outcomes. Fifth, model interpretation must be aggregated into economic feature groups. Sixth, the results must be exposed through a dashboard artifact while preserving caveats.",
        "The object of the research is the financial condition and forward outcome profile of U.S. public companies observed through structured filings and market-context variables. The subject of the research is the observable relationship between firm-level fundamentals, financial ratios, deterioration patterns, industry context, macro regimes, global-event windows, and future labels for strict distress, broader financial pressure, and resilience. This wording is important because the thesis does not claim to observe all strategic resources or internal managerial capabilities. It studies observable financial and contextual proxies that can be reproduced from public data.",
        "The research methods combine theoretical and empirical techniques. The theoretical part uses critical literature analysis and synthesis to connect strategic-management framing with financial-distress prediction and business analytics. The empirical part uses data engineering, panel construction, forward-label design, temporal validation, supervised classification, rare-event metrics, calibration and ranking diagnostics, feature-group interpretation, and audit controls. The artifact part uses dashboard design principles to make the empirical system inspectable. This combination is appropriate for a business analytics and big data systems program because it links methodological rigor with a practical analytical tool.",
        "The scientific novelty is concentrated in integration and discipline rather than in inventing a new statistical estimator. The thesis integrates filing-date-aligned SEC accounting data, FRED macro variables, curated global-event windows, target hierarchy, leakage controls, selected-fact provenance, temporal validation, feature-group interpretation, and dashboard exploration. Each component has precedents in prior work, but their audited combination within one empirical system addresses a practical gap in firm-level market-shift analytics. The novelty is therefore architectural and methodological, not algorithmic.",
        "The practical significance is that the framework can support structured monitoring and explanation. A financial analyst can compare strict distress risk with broader financial pressure and positive resilience outcomes. A manager can inspect whether model signal is mainly linked to profitability, leverage, cash-flow quality, deterioration, sector context, or market regime. A researcher can reuse the target hierarchy to test alternative model families. A dashboard user can move from aggregate sector patterns to firm-level trajectories without losing sight of target definitions and missingness caveats.",
        "The structure of the thesis follows the research tasks. The literature review establishes why success and failure must be interpreted through strategic, financial, macro, and analytical lenses. The methodology chapter explains the empirical panel, target construction, validation design, model families, audit controls, and dashboard artifact. The results chapter reports target counts, model metrics, calibration and ranking evidence, feature-group interpretation, and dashboard findings. The discussion chapter interprets implications, limitations, and future extensions. The conclusion returns to the research problem and states the contribution in bounded terms.",
    ],
    "2. Literature Review": [
        "The resource-based view and dynamic-capabilities literature are useful because they shift attention away from a purely mechanical view of failure. A company does not become resilient simply because one ratio is high; it survives market shifts when its resources, operating routines, financing structure, and ability to adjust remain adequate for the environment. At the same time, these theories are not directly measurable through SEC accounting tags. Their role in this thesis is therefore framing rather than direct variable definition. They justify the question of heterogeneous firm outcomes, while the empirical model measures financial manifestations of that heterogeneity.",
        "Organizational ecology provides a contrasting interpretation. Where dynamic-capabilities theory emphasizes adaptation, ecology emphasizes inertia, selection, and environmental fit. This contrast is useful for interpreting results because the model may identify deterioration before a formal event without observing whether managers attempted to adapt. A firm with weakening margins, increasing leverage, and repeated losses may be experiencing strategic inertia, poor market fit, or temporary shock exposure. The empirical framework can detect the pattern but cannot assign a single organizational cause. This reinforces the need to write findings as associations rather than causal explanations.",
        "Classical bankruptcy-prediction papers remain important because they establish the continuing relevance of accounting ratios. Beaver showed that individual ratios can contain failure information; Altman showed that ratios can be combined; Ohlson introduced probability logic; Zmijewski emphasized estimation and sample-selection problems. The present thesis differs from these works in data scale, temporal design, and model family, but it inherits their central insight: financial statements contain structured signals about future risk. The difference is that the signal is now embedded in a broader panel with macro context, industry metadata, and multiple outcomes.",
        "The hazard-model tradition is especially relevant because it criticizes static designs. A firm-period panel asks a different question from a one-time failed-versus-nonfailed comparison. The same company can appear many times with changing financial condition, changing macro context, and changing proximity to future events. This makes temporal ordering essential. Shumway's critique supports the use of repeated observations and forward-time validation even when the implemented model is not a formal hazard model. The thesis adopts the panel logic while using supervised classification models that fit the dashboard and feature-importance workflow.",
        "Market-based distress models create an important boundary. Structural default models and distance-to-default approaches use equity values and volatility to infer default risk from market information. They are conceptually powerful, but they require data and assumptions not implemented in the production panel. This thesis therefore treats them as future extensions and methodological comparators. The distinction prevents overclaiming. The study is an accounting, macro, event, and dashboard artifact; it is not a full market-microstructure or option-pricing default model.",
        "Modern machine-learning studies show that nonlinear models often improve classification in financial risk settings, yet they also create interpretation problems. Random forests and gradient boosting can capture interactions between profitability, leverage, liquidity, sector context, and macro variables. However, a high-performing model can still be invalid if the label is poorly defined or if future information leaks into features. This is why the thesis gives equal attention to target hierarchy and audit controls. The algorithm is only one part of the empirical design.",
        "The credit-scoring literature is related but not identical to firm failure analysis. Credit scoring often deals with borrower-level default or delinquency under institutional data systems, while this thesis studies public-company firm-period observations using SEC and macro data. The connection is methodological: both settings require model comparison, imbalanced-class metrics, calibration, ranking, and interpretability. The difference is substantive: firm success and failure across industries include strategic and accounting dynamics that extend beyond individual credit classification.",
        "Recent systematic reviews are valuable because they show that bankruptcy prediction remains active after the pandemic period and that model evaluation practices still vary widely. Some studies emphasize algorithmic novelty, while others focus on feature engineering, class imbalance, or interpretability. The thesis uses this literature to argue that a stronger contribution comes from validation and transparency rather than from simply adding a more complex model. In a high-stakes financial setting, a transparent random forest with correct timestamps may be more defensible than a complex model trained under weak temporal assumptions.",
        "The macro-finance literature provides the theoretical basis for treating market shifts as more than background noise. Credit spreads, financial conditions, interest rates, inflation, oil prices, dollar strength, and lending standards can affect financing costs, demand, margins, and refinancing risk. However, macro variables operate through firm exposures and balance-sheet conditions. A high-rate regime may be manageable for a cash-rich firm with strong margins but dangerous for a leveraged firm with declining revenue. This interaction logic explains why macro context belongs in the panel even when firm features dominate model importance.",
        "Global-event context adds a discrete-shock layer that continuous macro series do not fully capture. Pandemic windows, banking stress, oil shocks, geopolitical events, and supply-chain disruptions can cluster changes in firm outcomes and sector behavior. Still, a curated event calendar is not a complete news database. It is a transparent historical context layer. This is a methodological compromise: it improves interpretability and dashboard storytelling without opening a separate large-scale natural-language-processing project that would require its own validation.",
        "The rare-event evaluation literature is central to interpreting strict legal distress. ROC-AUC can look acceptable while precision remains low because positive cases are scarce. Precision-recall metrics are therefore more informative for the strict target. The thesis reports PR-AUC and F1 alongside test positives and threshold behavior to avoid misleading conclusions. This choice is not only technical. It affects the substantive story: strict distress is a benchmark and ranking problem, while broader pressure and resilience are richer factor-analysis outcomes.",
        "The missing-data literature is also more than a technical footnote. In accounting panels, missingness can reflect reporting choices, non-applicability, taxonomy variation, business model differences, dimensional reporting, or derivation limits. Treating every missing value as zero would create artificial information. Treating every missing value as unchanged would create another unsupported assumption. The thesis therefore preserves nulls in the empirical panel and limits imputation to the model-estimation pipeline. This approach aligns statistical caution with accounting reality.",
        "Explainable artificial intelligence literature supports the demand for interpretation but warns against overinterpretation. SHAP, LIME, permutation importance, and feature importance can help explain model behavior, but they do not prove causal mechanisms. The thesis therefore emphasizes economic feature groups. Group-level interpretation is less fragile than isolated feature rankings and more aligned with management concepts such as profitability, leverage, liquidity, cash-flow quality, deterioration, and industry exposure. This choice makes the explanation more communicable without hiding model complexity.",
        "Design-science literature is appropriate because the dashboard is not merely a visualization of results. It is a research artifact that embodies target definitions, data coverage, model outputs, and limitations. A dashboard can make a predictive model more accountable when it exposes denominators, missingness, target timelines, and caveats. It can also mislead if it presents scores without context. The thesis treats dashboard design as part of the methodology because the practical value of predictive analytics depends on how outputs are inspected and used.",
        "International and emerging-market studies are useful for limiting the scope of claims. Predictors that work in U.S. public-company filings may not transfer to private firms, emerging-market issuers, or jurisdictions with different accounting standards and legal procedures. Russia-specific and emerging-market research supports this caution. The thesis can therefore motivate future international extension while keeping the empirical claims tied to the U.S. SEC/FRED panel. This boundary is essential for scientific validity.",
        "The critical synthesis is that prior work provides all the components but not the exact integrated artifact. Classical models justify ratios, dynamic models justify panels, macro-finance justifies regime context, machine learning justifies nonlinear classifiers, rare-event literature justifies PR-AUC, XAI justifies interpretation, missing-data literature justifies null policy, and design science justifies the dashboard. The contribution is to combine these components into one audited framework and to report results under a target hierarchy that distinguishes legal distress, broader pressure, and resilience.",
    ],
    "3. Research Methodology": [
        "Panel construction begins from a clear observation rule. Each row represents one firm, one accounting period, and one information date. The information date is the SEC filing date rather than the fiscal period end date. This distinction is crucial because the model is intended to represent what could be known when the filing became available. The accounting period describes the economic interval, but the filing date determines the prediction information set. This design reduces look-ahead bias and aligns the study with early-warning logic.",
        "The SEC accounting layer is based on standardized financial statement data, but standardization does not eliminate all accounting complexity. Companies use different tags, some concepts are reported cumulatively, some fields are not separately reported by all industries, and amended filings can duplicate information dates. The methodology therefore uses a conservative concept map, current-period filtering, cumulative-flow derivation where justified, and deterministic duplicate handling. Risky shortcuts, such as deriving total liabilities silently from assets minus equity, are avoided because they would weaken interpretability and auditability.",
        "Flow variables require special treatment because quarterly financial statements can report either single-period values or year-to-date cumulative values. When a current-period flow is available, it is used directly. When only a cumulative flow is available, the single-period value can be derived by subtracting the prior cumulative value within the same firm, fiscal year, and variable. This derivation is an accounting transformation from reported facts rather than invented data. If the prior cumulative value is unavailable, the single-period value remains missing.",
        "The selected-fact provenance audit is a key methodological control. It checks whether the values chosen for the panel correspond to the intended accounting period and whether selected values match the panel values within tolerance. The current provenance output contains 700,844 selected facts, zero selected facts with a mismatched period date, and zero selected-value mismatches above tolerance. This evidence supports the use of the SEC-derived variables while preserving the caveat that taxonomy harmonization is conservative rather than perfect.",
        "Identifier handling is another methodological issue. Tickers can change, become delisted, or represent aliases, while CIKs provide a more stable SEC identifier. The empirical panel therefore uses CIK as the stable company key and treats ticker or display identifiers as readability fields. The PCG and PGNPQ case illustrates the importance of this approach: a bankruptcy-related source ticker can refer to a historical alias of the same CIK. The final panel keeps one CIK history while preserving alias metadata for event matching.",
        "The macroeconomic layer is aligned to prediction dates. This means that macro variables and regime indicators are attached according to information available at or before the filing-date observation. The macro set includes rates, Treasury spreads, volatility, credit spreads, financial conditions, inflation, oil prices, dollar strength, labor-market measures, retail sales, housing starts, commercial and industrial loans, and lending standards. Derived year-over-year changes convert selected level series into comparable growth or shock measures.",
        "Regime indicators translate continuous macro variables into interpretable states. High-rate, market-stress, tight-financial-conditions, credit-spread-stress, yield-curve-inversion, inflation-pressure, oil-shock, strong-dollar, demand-slowdown, credit-tightening, and crisis-regime flags are not meant to be perfect macroeconomic classifications. Their value is interpretability. They allow the dashboard and analysis to ask whether firm outcomes differ during recognizable market environments.",
        "The global-event layer is deliberately curated. It includes event windows for financial crises, sovereign debt episodes, natural disasters, commodity oil shocks, financial-market events, political-policy shifts, trade-policy episodes, pandemic conditions, supply-chain stress, monetary-inflation events, geopolitical war, and banking stress. The methodology treats these events as historical context fields. They are not used to claim that the model forecasts events or that all relevant events are captured.",
        "Feature families are separated to preserve interpretation. Accounting fundamentals are directly mapped or derived from SEC facts. Financial ratios standardize firm condition relative to size or revenue. Firm trend and deterioration features use prior observations to capture direction of change. Macro variables and regimes describe the external environment. Industry metadata enables cross-sector comparison. This family structure supports both modeling and explanation because feature importance can be aggregated into economically meaningful groups.",
        "Target construction is the most important methodological choice after timestamp alignment. The strict legal distress target uses formal event dates within a forward horizon. The broader failure-pressure target combines formal distress with repeated future profitability weakness and additional stress signals. The success/resilience target identifies future observations with positive net income, positive return on assets, acceptable leverage, and no formal distress. Secondary outcomes capture relative resilience, stress resilience, recovery, and cash-flow-supported success. Each target has a distinct denominator and interpretation.",
        "The strict target is intentionally conservative. It is positive only when a formal distress event occurs after the prediction date and within the defined horizon. Because the event is legal and rare, the positive count is small relative to the full panel. This makes it inappropriate as the sole headline target for a thesis about success and failure factors. It remains valuable because it anchors the analysis to a clean event definition and tests whether the system can rank observations before formal distress.",
        "The broader failure-pressure target is designed for factor analysis. It captures economically serious deterioration even when no formal bankruptcy event is observed. The target requires repeated future weakness and additional stress logic rather than labeling any temporary negative quarter as failure. This conservative design supports interpretation: positive cases represent sustained financial pressure, not arbitrary underperformance. The target therefore bridges legal distress and broader business deterioration.",
        "The success/resilience target is missing-aware because future observations are required to determine whether a firm remains profitable and acceptably leveraged. If too few future observations or components are available, the label remains missing rather than becoming a negative case. This policy avoids punishing firms for unavailable future data. It also explains why the success/resilience target has fewer known rows than strict distress. The smaller denominator reflects a more careful label definition.",
        "Model rows are filtered to remove post-event rows from primary forward-looking training. Post-event observations remain useful for historical dashboard analysis, but they should not be used as ordinary pre-event observations in a model designed to predict future distress. This distinction protects the model from learning after-the-fact patterns and supports a cleaner early-warning interpretation. The post-event audit confirms that the handling of these rows is controlled.",
        "Temporal validation uses prediction years to separate training, validation, and test periods. Training covers earlier observations, validation supports model and threshold selection, and the test period evaluates forward-time generalization. This design is more realistic than a random row split because firm conditions and macro regimes change over time. Random splits could place nearby observations from the same firm and market period on both sides of the split, producing an overly optimistic estimate.",
        "The baseline model families are intentionally interpretable and reproducible. Logistic regression provides a linear reference point. Random forest captures nonlinear interactions and is often strong for tabular data. Gradient boosting provides another nonlinear tree-based benchmark. Controlled tuning and advanced boosted-tree checks add robustness evidence, but they do not replace the central interpretation. The thesis favors a stable and auditable modeling design over algorithmic novelty for its own sake.",
        "The preprocessing pipeline handles mixed data types. Numeric features are imputed within the training pipeline using medians learned from the training data. Categorical features are encoded in a way compatible with the model families. Missingness indicators can be created by the preprocessing pipeline, but they are treated as engineering features. They may improve prediction, yet they are excluded from economic feature-group interpretation so that reporting gaps are not misrepresented as financial mechanisms.",
        "Evaluation metrics are selected according to target prevalence. ROC-AUC remains useful for ranking across thresholds, but PR-AUC is more informative when positives are rare. Precision, recall, F1, confusion matrices, and top-rank precision show whether alerts are practically meaningful. For resilience targets with higher positive shares, threshold interpretation differs because the base rate is already high. The thesis therefore reports metrics with denominators and test positives rather than comparing raw scores across targets without context.",
        "Calibration and ranking diagnostics add a practical layer beyond classification metrics. Calibration assesses whether predicted probabilities align with observed rates, while ranking diagnostics ask whether the top-scored observations are enriched with positives. For decision support, ranking can be more relevant than a single probability threshold. A risk analyst may inspect the top 5% or 10% of firms rather than rely on an automatic binary alert. This is why top-rank precision and lift are included in the empirical interpretation.",
        "Audit controls are treated as part of methodology rather than as administrative checks. Prediction timestamp, distress-event-date, leakage, macro-lag, global-event-timing, and post-event audits directly affect scientific validity. Feature blacklist checks protect against target leakage. Model-output consistency checks protect against reporting stale or mismatched results. Package and schema checks protect reproducibility. These controls collectively support the claim that the empirical system is ready for thesis interpretation with disclosed caveats.",
        "The dashboard artifact is designed to expose the same logic used in the methodology. It separates target composition, data coverage, model metrics, feature groups, firm trajectories, macro comparison, and artifact caveats. It does not hide denominator differences or missingness. It also separates target timelines from financial metrics so that binary labels are not visually confused with continuous accounting variables. The dashboard therefore supports practical interpretation and serves as a decision-support demonstration.",
        "The methodology has deliberate exclusions. Survival models, market-price distance-to-default models, full news-based NLP, ESG ratings, HR variables, patents, customer sentiment, and international accounting panels are not included in the production empirical system. These exclusions are not omissions of convenience. Each would require additional data rights, timestamp rules, cleaning, validation, and limitations. The thesis keeps the empirical core narrower so that the implemented claims remain defensible.",
    ],
    "4. Results": [
        "The first result is structural: the empirical system forms a consistent production panel. The panel includes 31,702 rows and 190 columns after the validated secondary target promotion. It covers 540 SEC CIKs and 540 ticker or display identifiers. The period-date range runs from 2009-03-31 to 2026-02-28, and the prediction-date range runs from 2009-04-15 to 2026-03-31. Processed and package-ready copies match exactly. This result matters because every subsequent model and dashboard output depends on a stable observation base.",
        "The second result is that target prevalence differs sharply across the hierarchy. Strict legal distress has 207 positives in 31,702 known rows. Broader failure pressure has 3,003 positives among 29,805 known rows. Success/resilience has 12,519 positives among 20,152 known rows. The four secondary targets have their own denominators and positive counts. These differences confirm that strict distress, broader pressure, and resilience are not interchangeable labels. They represent different empirical questions.",
        "Strict legal distress should be read as a rare-event benchmark. Its baseline gradient-boosting test PR-AUC is 0.104 and F1 is 0.182 with 96 test positives. These metrics are modest, but they are consistent with the difficulty of predicting rare formal events across a broad public-company panel. The model is not presented as a strong bankruptcy oracle. Its role is to show how a clean event target behaves under temporal validation and to anchor the broader target hierarchy.",
        "Broader failure pressure is the main failure-factor result. The baseline random forest reaches test PR-AUC of 0.758 and F1 of 0.725 with 575 test positives. This difference from strict distress is expected. Broader pressure is more common, more directly linked to accounting deterioration, and better suited for feature interpretation. The result supports the thesis decision to use failure pressure as the primary target for analyzing failure factors while keeping strict distress as a legal benchmark.",
        "Success/resilience is the strongest primary target. The baseline random forest reaches test PR-AUC of 0.975 and F1 of 0.938 with 2,648 test positives. The high score does not mean that the model has discovered the full theory of organizational success. It means that under the defined financial-resilience label, current and recent financial condition are strongly associated with future positive outcomes. The target is operational and financial, not a complete strategic resilience construct.",
        "The validated secondary outcomes extend the results beyond a simple positive/negative contrast. Industry-relative resilience has test PR-AUC of 0.855 and F1 of 0.787. Stress resilience has test PR-AUC of 0.974 and F1 of 0.937. Recovery has test PR-AUC of 0.977 and F1 of 0.915. Quality-success-cashflow has test PR-AUC of 0.953 and F1 of 0.920. These results support dimensional analysis of resilience, recovery, and success quality, but they do not replace the primary target hierarchy.",
        "Calibration diagnostics provide a probability-quality perspective. The selected validation-calibrated models show strong PR-AUC for broader pressure and resilience outcomes and lower PR-AUC for strict distress. For broader pressure, the calibrated test PR-AUC is 0.739 with a Brier score of 0.042. For success/resilience, calibrated test PR-AUC is 0.972 with a Brier score of 0.062. For strict distress, calibrated test PR-AUC remains near 0.100, reinforcing the rare-event interpretation.",
        "Ranking diagnostics show practical enrichment. In the strict distress test period, the top 10% of ranked observations captures 44.8% of positives with precision above the base rate. For broader failure pressure, the top 10% reaches precision of about 0.739 and recall of about 0.737, with lift above seven times the base rate. For success/resilience, top-ranked observations have very high precision, but the lift is smaller because the base rate is already high. These results support triage and monitoring use cases.",
        "Economic feature-group interpretation confirms the central factor result. For strict distress under gradient boosting, accounting fundamentals, firm trend/deterioration, and financial ratios account for almost all economic importance after excluding missingness indicators. For broader pressure under random forest, accounting fundamentals, trend/deterioration, and financial ratios are again dominant. For success/resilience under random forest, financial ratios are strongest, followed by trend/deterioration and accounting fundamentals. This pattern is stable with the thesis claim that firm-level financial condition drives the core empirical signal.",
        "Macro conditions and industry variables still have analytical value even when they are not dominant. Their lower feature-group share in the main tree-based models does not make them irrelevant. Macro and event variables provide the market-shift context required by the research question and allow comparisons across high-rate, inflation-pressure, market-stress, oil-shock, credit-tightening, and crisis environments. They are especially useful for dashboard filters and sector summaries, where context matters even if it does not dominate prediction.",
        "The ablation evidence is consistent with this interpretation. Firm-only and firm-plus-context feature sets both provide useful signal, but firm fundamentals and deterioration remain the empirical core. Event-aware specifications do not overturn the conclusion. The correct reading is that external conditions shape the environment in which firms operate, while the observable financial state of the firm remains the strongest direct predictor of the modeled outcomes.",
        "The data-quality results are themselves part of the empirical findings. All P0 audits pass. Selected SEC fact provenance shows current-period alignment and value consistency. Same-information-date amendment duplicates are resolved. Non-positive total-assets rows are eliminated through source-aware nulling, and remaining negative revenue observations are source-reported rather than unresolved mapping errors. These controls strengthen confidence that model results are not artifacts of obvious data leakage or unresolved period mismatch.",
        "The dashboard result is practical rather than purely statistical. The artifact exposes the composition of the panel, target coverage, model metrics, feature groups, firm-level histories, macro comparison, target timelines, and caveats. It makes the target hierarchy visible and prevents users from treating strict distress, broader pressure, and success/resilience as the same outcome. It also allows sector and event-window comparisons, which supports the cross-industry and market-shift framing of the thesis.",
        "The combined result is a coherent empirical story. Strict legal distress is rare and difficult, broader failure pressure is informative for failure-factor analysis, success/resilience is a strong positive counterpart, secondary outcomes add useful dimensions, and firm fundamentals, ratios, and deterioration dominate the interpretation layer. Macro regimes and global events contextualize rather than replace firm analysis. This story is more nuanced than a single accuracy table and better aligned with the research problem.",
    ],
    "5. Discussion": [
        "The main theoretical implication is that predictive analytics can operationalize parts of the success and failure debate without reducing the debate to a single ratio or a single legal event. Strategic theories describe resources, capabilities, inertia, and adaptation, while financial models observe their accounting traces imperfectly. The thesis bridges these levels by treating financial pressure and resilience as observable forward-looking proxies. This does not replace strategic theory, but it gives a structured empirical way to compare firms across industries and regimes.",
        "The first managerial implication is the importance of denominator awareness. A high success/resilience metric and a low strict-distress metric are not competing claims about the same target. They answer different questions with different base rates. Managers and analysts should therefore choose the target that matches the decision problem. If the task is legal-event warning, strict distress is relevant. If the task is early detection of serious operating and balance pressure, the broader target is more informative. If the task is identifying durable positive financial condition, success/resilience is appropriate.",
        "The second managerial implication is that factor groups are more actionable than raw model scores. A score without explanation can support ranking but not diagnosis. A feature-group view can show whether the model signal is mainly linked to leverage, profitability, liquidity, cash flow, deterioration, or sector context. This distinction matters because different feature groups imply different follow-up questions. A leverage signal may lead to refinancing analysis, while a deterioration signal may lead to revenue, margin, or operational review.",
        "The third managerial implication is that market shifts should be interpreted jointly with firm condition. A high-rate regime or oil shock does not affect all firms equally. The same macro state can be manageable for one firm and dangerous for another depending on margins, leverage, liquidity, sector exposure, and recent trends. The dashboard is useful because it allows analysts to inspect this joint context rather than attributing distress mechanically to the macro environment.",
        "For investors, the framework supports screening and prioritization. Strict distress rankings identify rare high-risk observations, while broader pressure rankings identify financially weak observations that may warrant deeper review. Success and secondary resilience rankings can support searches for stable or improving firms. However, the framework is not a buy-sell recommendation system. It lacks market-price valuation, portfolio constraints, transaction costs, and investor-specific objectives. Its value is analytical triage and explanation.",
        "For researchers, the target hierarchy is the most transferable part of the contribution. A similar hierarchy could be applied to other countries, private-company datasets, banking portfolios, or sector-specific studies, provided that outcome labels and timestamp controls are rebuilt for the new context. The empirical numbers in this thesis should not be copied across jurisdictions, but the design principle of separating clean rare events from broader pressure and positive resilience can travel.",
        "The first limitation is jurisdictional scope. SEC filings provide a rich and reproducible public-company data source, but they reflect U.S. reporting rules, public-company disclosure incentives, and U.S. market institutions. Private companies, emerging-market firms, and firms under different accounting standards may show different reporting patterns and distress dynamics. International extension would require new accounting harmonization, legal-event definitions, and macro-data timing controls.",
        "The second limitation is target validity. Strict legal distress is sourced and clean for the current formal rows, but it remains a rare event and not a complete universal legal database. Broader failure pressure is economically meaningful but not legal bankruptcy. Success/resilience is a financial proxy rather than total organizational resilience. Secondary targets capture dimensions of resilience and recovery, but they do not prove turnaround mechanisms. These distinctions must remain visible in any interpretation.",
        "The third limitation is that feature importance is not causality. Accounting fundamentals and ratios are economically meaningful, but the model estimates predictive association under the chosen validation design. A high feature-group contribution does not prove that changing one variable would change the outcome. Causal inference would require a different research design, identification strategy, and possibly intervention or quasi-experimental evidence. The thesis therefore uses the language of association, signal, and interpretation.",
        "The fourth limitation is missingness. Null values are preserved and diagnosed, but missingness remains material for some accounting concepts and targets. High missingness in fields such as R&D expense or gross profit can reflect sector differences, reporting conventions, or taxonomy limitations. The model can handle missing values through imputation and indicators, but economic interpretation must focus on substantive observed variables. This limitation is especially important for success and secondary targets because future components are not always observable.",
        "The fifth limitation concerns the dashboard artifact. The dashboard demonstrates decision-support logic and exposes empirical outputs, but it is not a governed production system. It does not include live SEC ingestion, operational monitoring, user-access controls, model drift detection, or formal user testing. These would be necessary for enterprise deployment. Within the thesis, the dashboard should be evaluated as a research artifact and exploratory interface.",
        "Future research can extend the framework in several directions. Market-price variables and distance-to-default measures would connect the accounting panel with structural credit-risk literature. Textual disclosures and earnings-call language could add forward-looking qualitative signals if processed with careful timestamp controls. International panels could test transferability across legal and accounting systems. User studies could evaluate whether the dashboard improves analyst understanding and decision quality. Each extension should preserve the core discipline of timestamp alignment, target clarity, and auditability.",
    ],
    "Conclusion": [
        "The research tasks are therefore addressed in sequence. The literature review establishes that firm success and failure under market shifts require strategic, financial, macro, methodological, and interpretability perspectives. The data task is addressed by constructing a filing-date-aligned SEC/FRED/global-event firm-period panel. The target task is addressed by separating strict distress, broader pressure, success/resilience, and secondary outcomes. The modeling task is addressed through temporal validation and appropriate metrics. The interpretation task is addressed through feature groups and dashboard views.",
        "The empirical answer to the research problem is that observable success and failure factors can be studied most defensibly through a target hierarchy and audited temporal panel rather than through a single binary label. Strict legal distress remains rare and difficult, broader pressure provides the main failure-factor lens, and success/resilience provides the positive counterpart. Secondary outcomes add useful dimensions but do not replace the primary structure. This design makes the thesis more accurate and more practically useful.",
        "The substantive answer is that firm fundamentals, financial ratios, and deterioration features carry the main predictive signal in the current empirical system. Macro regimes and global events are still important because they describe the conditions in which firms operate and enable sector comparison, but they do not dominate the central feature-group interpretation. This finding is consistent with the view that market shifts matter through their interaction with firm condition rather than as standalone explanations.",
        "The final contribution is a reproducible business-analytics artifact with academic and practical value. It combines data construction, target design, temporal validation, audit controls, model interpretation, and dashboard exploration. Its claims are intentionally bounded: the framework supports predictive association, factor interpretation, and decision support; it does not prove causality or provide a live bankruptcy oracle. Within those boundaries, the thesis provides a coherent method for navigating market shifts through predictive insights into success and failure factors across industries.",
    ],
}


def maybe_expand_body(main_body: list[tuple[str, int, list[str]]], target_chars: int = 98000) -> None:
    for title, _, paras in main_body:
        paras.extend(ACADEMIC_EXTENSIONS.get(title, []))


def add_front_matter(doc: Document) -> None:
    add_heading(doc, "Table of Contents", 1)
    add_toc_field(doc)
    doc.add_page_break()

    add_heading(doc, "Abstract", 1)
    add_para(
        doc,
        "This thesis develops an applied predictive-analytics framework for studying company success and failure factors across industries under market shifts. The empirical artifact combines SEC Financial Statement Data Sets, Federal Reserve Economic Data, curated global-event context, industry metadata, and source-verified distress labels into one firm-period panel. The production panel contains 31,702 rows, 190 columns, 540 SEC CIKs, and 540 ticker or display identifiers. Prediction timestamps are anchored to SEC filing dates, so each observation is modeled from the date when the financial statement became publicly available. The study separates strict legal distress, broader conservative failure pressure, success/resilience, and four validated secondary production outcomes rather than forcing the empirical problem into one binary label.",
    )
    add_para(
        doc,
        "The results show that strict legal distress is a clean but rare benchmark, broader financial pressure is more suitable for failure-factor analysis, and success/resilience is a strong positive counterpart under missing-aware future-label logic. Feature-group interpretation indicates that accounting fundamentals, financial ratios, and firm deterioration carry the main economic signal, while macro regimes and global-event windows provide market-shift context and dashboard segmentation. The thesis contributes an audited panel, target hierarchy, temporal validation design, leakage and provenance controls, interpretable model outputs, and a dashboard artifact for historical exploration and decision support. The evidence is interpreted as predictive association and decision-support insight, not as proof of causal mechanisms.",
    )
    add_para(doc, "Keywords: predictive analytics; financial distress; firm resilience; SEC Financial Statement Data Sets; FRED; temporal validation; machine learning; dashboard; market shifts.")

    add_heading(doc, "List of Abbreviations", 1)
    abbreviations = [
        ("CIK", "Central Index Key, the stable company identifier used by the U.S. Securities and Exchange Commission."),
        ("EDGAR", "Electronic Data Gathering, Analysis, and Retrieval system."),
        ("FRED", "Federal Reserve Economic Data."),
        ("ML", "Machine learning."),
        ("NFCI", "National Financial Conditions Index."),
        ("P0 audit", "Blocking integrity audit used to check prediction timestamp, leakage, macro timing, global-event timing, distress-event dates, and post-event row handling."),
        ("PR-AUC", "Precision-recall area under the curve."),
        ("ROC-AUC", "Receiver operating characteristic area under the curve."),
        ("SEC", "U.S. Securities and Exchange Commission."),
        ("SHAP", "SHapley Additive exPlanations."),
        ("VIX", "CBOE Volatility Index."),
    ]
    for key, value in abbreviations:
        add_bullet(doc, f"{key}: {value}")
    doc.add_page_break()


def add_results_tables(doc: Document, section_title: str) -> None:
    if section_title == "1. Introduction":
        add_heading(doc, "1.1 Research Design Overview", 2)
        add_table(
            doc,
            ["Element", "Definition in the study"],
            [
                ["Research problem", "Identifying and interpreting observable success and failure factors under changing market conditions."],
                ["Research object", "U.S. public-company firm-period observations constructed from SEC filings and market-context data."],
                ["Research subject", "Relationships between accounting fundamentals, ratios, deterioration, industry context, macro regimes, global events, and forward outcomes."],
                ["Main empirical artifact", "An audited firm-period panel and decision-support dashboard."],
                ["Central boundary", "Predictive association and model interpretation are analyzed; causal proof is not claimed."],
            ],
            "author's synthesis from the empirical dataset and methodology.",
        )
    if section_title == "3. Research Methodology":
        add_heading(doc, "3.1 Data Sources Used in the Empirical Artifact", 2)
        add_table(
            doc,
            ["Source family", "Role in thesis", "Important limitation"],
            [
                ["SEC Financial Statement Data Sets", "Accounting fundamentals and filing dates.", "Requires concept mapping, period selection, and provenance audits."],
                ["FRED macro-financial series", "Rates, spreads, inflation, labor, demand, credit, and market context.", "Contextual variables, not firm-level causality by themselves."],
                ["Curated global-event calendar", "Historical market-shift windows for dashboard and sector/event comparison.", "Not a full live news/event database."],
                ["Distress-event date file", "Strict legal distress labels and forward event windows.", "Clean benchmark, but rare and not a universal legal database."],
                ["Industry and firm metadata", "Sector filters, CIK/ticker identifiers, SIC and dashboard context.", "Sampling metadata is excluded from predictive features when it would leak design logic."],
            ],
            "author's synthesis from the empirical dataset design.",
        )
        add_heading(doc, "3.2 Current Production Panel", 2)
        add_table(
            doc,
            ["Item", "Current value"],
            [
                ["Rows", "31,702"],
                ["Columns", "190"],
                ["SEC CIKs", "540"],
                ["Ticker/display identifiers", "540"],
                ["Period-date range", "2009-03-31 to 2026-02-28"],
                ["Prediction-date range", "2009-04-15 to 2026-03-31"],
                ["Prediction-date policy", "prediction_date = filed_date"],
            ],
            "author's calculations from the production panel.",
        )
        add_heading(doc, "3.3 Target Hierarchy and Label Counts", 2)
        add_table(
            doc,
            ["Target", "Known rows", "Positives", "Missing rows", "Role"],
            PRIMARY_TARGET_TABLE,
            "author's calculations from the production panel.",
        )
    if section_title == "4. Results":
        add_heading(doc, "4.1 Baseline Temporal Test Metrics", 2)
        add_table(
            doc,
            ["Target", "Best baseline model", "Test rows", "Test positives", "Test PR-AUC", "Test F1"],
            PRIMARY_METRICS_TABLE,
            "author's calculations from temporal model outputs.",
        )
        add_heading(doc, "4.2 P0 Audit Summary", 2)
        add_table(
            doc,
            ["Audit", "Status", "Interpretation"],
            AUDIT_TABLE,
            "author's audit summary from the empirical validation outputs.",
        )
        add_heading(doc, "4.3 Economic Feature-Group Interpretation", 2)
        add_table(
            doc,
            ["Target", "Model", "Top economic feature group", "Share excluding missingness"],
            FEATURE_GROUP_TABLE,
            "author's calculations from feature-group interpretation outputs.",
        )


def add_main_body(doc: Document) -> int:
    maybe_expand_body(MAIN_BODY)
    main_chars = 0
    for title, level, paragraphs in MAIN_BODY:
        add_heading(doc, title, level)
        add_results_tables(doc, title)
        for paragraph in paragraphs:
            add_para(doc, paragraph)
            main_chars += len(paragraph)
        if title not in {"Conclusion"}:
            doc.add_section(WD_SECTION.NEW_PAGE)
    return main_chars


def add_references(doc: Document) -> None:
    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "References", 1)
    refs = load_references()
    for idx, ref in enumerate(refs, start=1):
        p = doc.add_paragraph(style="List Number")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(ref)
        run.font.name = BODY_FONT
        run.font.size = BODY_SIZE


def add_appendices(doc: Document, main_chars: int) -> None:
    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "Appendices", 1)
    add_heading(doc, "Appendix A. Empirical Scope Summary", 2)
    add_table(
        doc,
        ["Element", "Scope in the thesis"],
        [
            ["Empirical setting", "U.S. public companies with standardized SEC financial statement data and macro-financial context."],
            ["Panel size", "31,702 firm-period observations, 190 columns, 540 SEC CIKs, and 540 ticker/display identifiers."],
            ["Prediction timestamp", "SEC filing date for every observation in the production panel."],
            ["Primary outcomes", "Strict legal distress, broader conservative failure pressure, and success/resilience."],
            ["Secondary outcomes", "Industry-relative resilience, stress resilience, recovery, and cash-flow-supported success."],
            ["Main interpretation boundary", "Predictive association and model interpretation are analyzed; causal mechanisms are not claimed."],
        ],
        "summary of the empirical scope used in the thesis.",
    )
    add_heading(doc, "Appendix B. Feature Families", 2)
    feature_rows = [
        ["Identifiers and timing", "adsh, cik, ticker, period_date, filed_date, prediction_date, form, AFS, SEC ZIP."],
        ["Accounting fundamentals", "Assets, liabilities, equity, revenue, operating income, net income, cash flow, capex."],
        ["Financial ratios", "ROA, leverage/assets, current ratio, cash/assets, margins, R&D intensity, receivables/assets."],
        ["Trends and deterioration", "Lagged ratios, four-observation changes, growth rates, prior negative-income counts."],
        ["Macro conditions", "Rates, spreads, VIX, NFCI, inflation, oil, dollar, labor, demand, housing, credit."],
        ["Regime indicators", "High-rate, market-stress, tight-financial-conditions, inflation-pressure, oil-shock, crisis, and related flags."],
        ["Global-event context", "Event counts, severity, event-type counts, and event names from curated event windows."],
        ["Industry metadata", "Sector, industry, SIC, SIC sector, exchange."],
    ]
    add_table(doc, ["Feature family", "Examples"], feature_rows, "summary of the empirical feature design.")

    add_heading(doc, "Appendix C. Thesis-Safe Claim Boundaries", 2)
    safe_claims = [
        "The thesis builds an audited SEC/FRED/global-event firm-period panel.",
        "Strict legal distress is treated as a clean but rare benchmark.",
        "Broader financial pressure is the main failure-factor analysis target.",
        "Success/resilience is a separate positive outcome.",
        "Validated secondary production outcomes support dimensional analysis.",
        "Feature importance describes model association, not causality.",
        "The dashboard is a historical decision-support artifact, not a live bankruptcy oracle.",
    ]
    for claim in safe_claims:
        add_bullet(doc, claim)

    add_heading(doc, "Appendix D. Dashboard Modes", 2)
    modes = [
        ["Overview", "Dataset and target composition, denominator differences, and coverage."],
        ["Model results", "Strict distress, broader failure pressure, success/resilience, and secondary target metrics."],
        ["Factor groups", "Economic feature-group interpretation excluding missingness indicators."],
        ["Firm explorer", "Latest filing, financial trajectory, macro context, and target timeline."],
        ["Macro/event comparison", "Market regimes and global-event windows for sector and historical context."],
        ["Artifact notes", "Scope, caveats, data lineage, and claim boundaries."],
    ]
    add_table(doc, ["Mode", "Purpose"], modes, "summary of the dashboard artifact design.")


def main() -> None:
    doc = Document()
    configure_document(doc)
    add_front_matter(doc)
    main_chars = add_main_body(doc)
    add_references(doc)
    add_appendices(doc, main_chars)
    doc.core_properties.title = "Navigating Market Shifts: Predictive Insights into Success and Failure Factors Across Industries"
    doc.core_properties.subject = "Master's thesis manuscript"
    doc.core_properties.author = "Artem Kozin"
    doc.core_properties.comments = "Master's thesis manuscript."
    doc.save(OUT)
    print(OUT)
    print(f"main_body_characters={main_chars}")
    print(f"references={len(load_references())}")


if __name__ == "__main__":
    main()
