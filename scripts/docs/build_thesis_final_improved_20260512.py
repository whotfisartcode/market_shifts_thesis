from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "thesis_final_improved_20260512.docx"

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(14)
SMALL_SIZE = Pt(11)
TABLE_SIZE = Pt(9)
APPENDIX_TABLE_SIZE = Pt(8)


def set_update_fields(doc: Document) -> None:
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def set_run_font(run, size=BODY_SIZE, bold: bool = False, italic: bool = False) -> None:
    run.font.name = BODY_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    for kind, text in [
        ("begin", None),
        ("instrText", "PAGE"),
        ("separate", None),
        ("text", "1"),
        ("end", None),
    ]:
        if kind == "instrText":
            elem = OxmlElement("w:instrText")
            elem.set(qn("xml:space"), "preserve")
            elem.text = text
        elif kind == "text":
            elem = OxmlElement("w:t")
            elem.text = text
        else:
            elem = OxmlElement("w:fldChar")
            elem.set(qn("w:fldCharType"), kind)
        run._r.append(elem)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.header_distance = Mm(12)
    section.footer_distance = Mm(12)
    section.different_first_page_header_footer = True
    add_page_number(section.footer.paragraphs[0])

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Inches(0.49)

    for name, size in [("Title", Pt(16)), ("Heading 1", Pt(14)), ("Heading 2", Pt(14)), ("Heading 3", Pt(14))]:
        style = styles[name]
        style.font.name = BODY_FONT
        style.font.size = size
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = BODY_FONT
        style.font.size = BODY_SIZE
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(3)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        set_run_font(run, BODY_SIZE, bold=True)


def add_para(doc: Document, text: str, *, first_line: bool = True, italic: bool = False, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Inches(0.49) if first_line else Inches(0)
    run = p.add_run(text)
    set_run_font(run, BODY_SIZE, bold=bold, italic=italic)


def add_small_para(doc: Document, text: str, *, italic: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    set_run_font(run, SMALL_SIZE, italic=italic)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    set_run_font(run, BODY_SIZE)


def add_numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    set_run_font(run, BODY_SIZE)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False, size=TABLE_SIZE) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    set_run_font(run, size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, number: str, title: str, headers: list[str], rows: list[list[str]], note: str, source: str, size=TABLE_SIZE) -> None:
    add_small_para(doc, f"{number}. {title}", italic=False)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_text(hdr[idx], header, bold=True, size=size)
        set_cell_shading(hdr[idx], "EDEDED")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, size=size)
    add_small_para(doc, f"Note: {note} Source: {source}", italic=True)


def add_toc_field(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and update field to refresh the table of contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep, placeholder, end])


def add_title_page(doc: Document) -> None:
    lines = [
        ("FEDERAL STATE AUTONOMOUS", 12, False),
        ("EDUCATIONAL INSTITUTION OF HIGHER EDUCATION", 12, False),
        ("NATIONAL RESEARCH UNIVERSITY", 12, False),
        ('"HIGHER SCHOOL OF ECONOMICS"', 12, True),
        ("Graduate School of Business", 12, False),
        ("", 12, False),
        ("Artem Kozin", 14, False),
        ("", 12, False),
        ("Navigating Market Shifts: Predictive Insights into Success and Failure Factors Across Industries", 16, True),
        ("", 12, False),
        ("MASTER'S THESIS", 14, True),
        ("", 12, False),
        ('Field of Study 38.04.05 "Business Informatics"', 12, False),
        ('Educational Programme "Business Analytics and Big Data Systems"', 12, False),
    ]
    for text, size, bold in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_after = Pt(4 if text else 10)
        run = p.add_run(text)
        set_run_font(run, Pt(size), bold=bold)

    for _ in range(5):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run("Supervisor\nArmen Beklaryan\nAssociate Professor")
    set_run_font(run, Pt(12))

    for _ in range(8):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run("Moscow, 2026")
    set_run_font(run, Pt(12))
    doc.add_page_break()


def add_landscape_section(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.left_margin = Mm(20)
    section.right_margin = Mm(15)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(15)
    section.footer.is_linked_to_previous = True


def add_figure(doc: Document, number: str, title: str, image_path: Path, caption: str) -> None:
    add_small_para(doc, f"{number}. {title}", italic=False)
    if image_path.exists():
        doc.add_picture(str(image_path), width=Inches(9.5))
    else:
        add_small_para(doc, f"[Missing image: {image_path}]", italic=True)
    add_small_para(doc, f"Caption: {caption}", italic=True)


def build_document() -> None:
    doc = Document()
    configure_document(doc)
    set_update_fields(doc)

    add_title_page(doc)

    add_heading(doc, "Table of Contents", 1)
    add_toc_field(doc)
    doc.add_page_break()

    add_heading(doc, "Abstract", 1)
    add_para(
        doc,
        "This thesis develops and evaluates an audited predictive-analytics framework for studying observable company success and failure factors across industries under changing market conditions. The empirical artifact combines SEC Financial Statement Data Sets, SEC filing metadata, FRED macro-financial variables, a curated global-event context layer, source-verified strict distress events, supervised temporal models, leakage audits, feature-family interpretation, and a dashboard artifact. The unit of analysis is a U.S. public-company firm-period observation dated by the SEC filing date, not by the fiscal period end date. This design makes the empirical question forward-looking while keeping the claim boundary clear: the thesis studies predictive association, not causality.",
    )
    add_para(
        doc,
        "The accepted production panel contains 31,702 rows, 190 columns, 540 SEC CIKs, and 540 ticker or display identifiers, with period dates from 2009-03-31 to 2026-02-28 and prediction dates from 2009-04-15 to 2026-03-31. The target hierarchy separates strict legal distress, broader conservative failure pressure, success/resilience, and four validated secondary outcomes. Results show that strict legal distress is a clean but rare benchmark, with baseline temporal test PR-AUC of 0.104 and F1 of 0.182. Broader failure pressure is the main failure-factor target, with random forest test PR-AUC of 0.758 and F1 of 0.725. Success/resilience is the strongest positive counterpart, with random forest test PR-AUC of 0.975 and F1 of 0.938. Feature-family interpretation indicates that accounting fundamentals, financial ratios, and firm trend or deterioration features provide the main economic signal, while macro regimes and global-event windows contextualize market shifts and dashboard exploration.",
    )
    add_para(
        doc,
        "Keywords: predictive analytics; financial distress; firm resilience; SEC Financial Statement Data Sets; FRED; temporal validation; machine learning; dashboard; market shifts.",
        first_line=False,
        italic=True,
    )

    add_heading(doc, "List of Abbreviations", 1)
    for item in [
        "CIK: Central Index Key, the stable company identifier used by the U.S. Securities and Exchange Commission.",
        "EDGAR: Electronic Data Gathering, Analysis, and Retrieval system.",
        "FRED: Federal Reserve Economic Data.",
        "ML: Machine learning.",
        "NFCI: National Financial Conditions Index.",
        "P0 audit: Blocking project-integrity audit covering prediction timestamps, event-date provenance, leakage, macro timing, global-event timing, and post-event row handling.",
        "PR-AUC: Precision-recall area under the curve.",
        "ROC-AUC: Receiver operating characteristic area under the curve.",
        "SEC: U.S. Securities and Exchange Commission.",
        "VIX: CBOE Volatility Index.",
        "XAI: Explainable artificial intelligence.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "1. Introduction", 1)
    for text in [
        "The ability to interpret firm success and failure under market shifts is a central problem for business analytics. Companies disclose large volumes of structured accounting data, while macro-financial conditions change through interest-rate cycles, credit-market stress, inflation shocks, oil-price movements, supply-chain disruptions, banking stress, and geopolitical events. Managers, analysts, investors, and regulators therefore face a practical question: which observable signals help separate resilient firms from firms moving toward serious pressure, and how can those signals be evaluated without confusing prediction with explanation or correlation with causation?",
        "The thesis title, Navigating Market Shifts: Predictive Insights into Success and Failure Factors Across Industries, requires a broader design than a conventional bankruptcy-prediction exercise. Formal legal distress is important, but it is rare and legally specific. Companies can experience repeated financial pressure without filing for bankruptcy, and companies can show financial resilience without proving overall organizational success in the strategic sense. The empirical design therefore uses a target hierarchy. Strict legal distress anchors the study to a clean event definition. Broader conservative failure pressure serves as the main failure-factor outcome. Success/resilience serves as a positive counterpart. Secondary outcomes add sector-relative resilience, stress resilience, recovery, and cash-flow-supported success.",
        "The relevance of the topic follows from both practice and literature. Corporate distress remains visible after the global financial crisis, the pandemic shock, the inflation and rate-hiking cycle, banking-stress episodes, and sector-specific disruptions. Recent bankruptcy and financial-stability reporting keeps distress analysis practically current (Federal Reserve Board, 2025; S&P Global Market Intelligence, 2025). At the same time, the academic literature shows that accounting ratios, profitability, leverage, liquidity, macro-financial conditions, and repeated firm observations can be informative for default or distress prediction, while also warning that model performance depends heavily on target definition, validation design, and sample construction.",
        "Classical research by Beaver (1966), Altman (1968), Ohlson (1980), and Zmijewski (1984) established that accounting information can contain failure signals. Later work by Shumway (2001), Duffie, Saita and Wang (2007), Campbell, Hilscher and Szilagyi (2008), Tinoco and Wilson (2013), and Barboza, Kimura and Altman (2017) expanded the field toward dynamic, macro-aware, and machine-learning designs. However, many empirical studies still leave practical gaps for a thesis in business analytics: a single distress label can hide important differences between legal failure, financial weakness, recovery, and resilience; random row splitting can overstate performance in time-ordered financial panels; and explanation outputs can be mistaken for causal mechanisms.",
        "The research problem can be formulated as follows: how can a reproducible, temporally valid, and interpretable predictive-analytics framework be used to identify and compare observable success and failure factors across industries under market-shift conditions?",
        "The goal of the thesis is to develop and evaluate an audited predictive-analytics framework for studying firm failure pressure, strict legal distress, and financial success/resilience across U.S. public companies and market environments. The framework is evaluated not only by model metrics, but also by target validity, timestamp discipline, leakage controls, missingness handling, feature-family interpretation, and the dashboard artifact through which results are explored.",
    ]:
        add_para(doc, text)

    add_heading(doc, "1.1 Research Questions", 2)
    for rq in [
        "RQ1 (Factor identification). Which economic feature groups are most strongly associated with strict legal distress, broader conservative failure pressure, and success/resilience across U.S. public companies?",
        "RQ2 (Target definition and model performance). How does target definition affect predictive performance and interpretation under temporal validation?",
        "RQ3 (Market-shift context). How do macro regimes and global-event windows contextualize firm outcomes across industries?",
        "RQ4 (Decision-support artifact). How can an audited dashboard artifact support transparent interpretation of predictive outputs without implying causal proof?",
    ]:
        add_numbered(doc, rq)

    add_heading(doc, "1.2 Research Design Overview", 2)
    add_table(
        doc,
        "Table 1",
        "Research design overview",
        ["Element", "Definition in the thesis"],
        [
            ["Research problem", "Identification and interpretation of observable success and failure factors under changing market conditions."],
            ["Research object", "U.S. public-company firm-period observations built from SEC filings and market-context data."],
            ["Research subject", "Predictive association between accounting fundamentals, ratios, deterioration, industry context, macro regimes, global-event windows, and future firm outcomes."],
            ["Main artifact", "Audited firm-period panel, temporal models, feature-family interpretation, and dashboard artifact."],
            ["Central boundary", "Predictive association and model behavior are analyzed; causal proof and live bankruptcy prediction are not claimed."],
        ],
        "The table defines the thesis at the level required for committee readability.",
        "Author's synthesis from the empirical design and project audit outputs.",
    )

    for text in [
        "The research tasks are to review literature on strategic adaptation, financial distress prediction, macro-financial market shifts, machine learning, rare-event evaluation, interpretability, missing data, and design science; to construct a firm-period panel from audited SEC, macro, event, and metadata sources; to formalize target definitions for strict distress, broader failure pressure, success/resilience, and secondary outcomes; to train and evaluate temporal baseline models; to audit leakage, timestamp alignment, event timing, and post-event rows; to interpret model behavior through economic feature groups; and to present results through a dashboard artifact with explicit caveats.",
        "The object of the research is the financial condition and forward outcome profile of public companies observed through structured filings and market-context variables. The subject of the research is the observable relationship between firm financial signals, contextual market-shift variables, and future labels representing legal distress, conservative failure pressure, and financial resilience. The main methods are literature synthesis, data engineering, feature construction, temporal train-validation-test splitting, supervised learning, PR-AUC and F1 evaluation, calibration and ranking diagnostics, leakage auditing, feature-family interpretation, and design-science artifact evaluation.",
        "The scientific novelty is concentrated in integration and discipline rather than in inventing a new bankruptcy algorithm. The thesis integrates a filing-date-aligned SEC/FRED/global-event panel, a target hierarchy, temporal validation, leakage audits, missing-aware label logic, feature-group interpretation, and a dashboard artifact into one reproducible business-analytics framework. The practical significance is that the framework supports transparent historical exploration of firm pressure and resilience, while making the limitations of the data, targets, and model outputs visible.",
        "The thesis is organized as follows. Chapter 1 presents the research problem, goal, questions, object, subject, methods, novelty, and practical significance. Chapter 2 reviews the literature and links each stream to a design decision. Chapter 3 describes the empirical methodology, including data sources, target definitions, temporal validation, leakage controls, and dashboard design. Chapter 4 reports panel integrity, target counts, model metrics, calibration and ranking diagnostics, feature-family interpretation, and dashboard evidence. Chapter 5 discusses implications, limitations, and future research. The conclusion answers RQ1 to RQ4 directly. Appendices provide target definitions, feature families, leakage controls, temporal split details, dashboard screenshots, and model-output support.",
    ]:
        add_para(doc, text)

    add_heading(doc, "2. Literature Review", 1)
    for text in [
        "The literature review is organized around six streams that justify the empirical design: strategic adaptation and resilience, classical distress prediction, dynamic and macro-financial default prediction, machine learning and rare-event evaluation, interpretability and decision support, and missing-data and data-integrity methods. The purpose of the review is not to describe every source separately, but to explain why this thesis uses a target hierarchy, a firm-period panel, filing-date timestamps, temporal validation, feature-family interpretation, and a dashboard artifact.",
        "Strategic-management theory motivates the broader language of success and failure. Barney (1991) argues that performance differences can be linked to resources when those resources are valuable, rare, difficult to imitate, and organizationally embedded. Teece, Pisano and Shuen (1997) shift the focus toward dynamic capabilities, which is important in market-shift settings because firms must adapt to changing environments. Hannan and Freeman (1984) provide a contrasting organizational-ecology view in which inertia and selection also shape survival. These theories justify the thesis title but also create a measurement boundary: SEC financial statements do not directly observe capabilities, culture, managerial quality, or strategic choices. The thesis therefore measures financial resilience, not total organizational success.",
        "The resilience literature strengthens this boundary. Linnenluecke (2017) and Duchek (2020) show that organizational resilience is a broad construct involving preparation, adaptation, response, and learning. Hepfer and Lawrence (2022) emphasize heterogeneity across forms of resilience. This means that a financial success/resilience target should be presented as a measurable counterpart, not as a complete theory of resilient organizations. In this thesis, success/resilience means future profitability, positive ROA, acceptable leverage, and absence of formal distress under a missing-aware rule.",
        "Classical distress prediction remains foundational because it establishes the economic relevance of accounting variables. Beaver (1966) demonstrated that financial ratios can separate failed and non-failed firms before failure. Altman (1968) combined ratios into a discriminant model, while Ohlson (1980) introduced probabilistic bankruptcy prediction and Zmijewski (1984) highlighted methodological concerns in distress modeling. Taffler (1983), Platt and Platt (1994), and Altman and Narayanan (1997) further show that solvency and failure classification are long-standing research problems rather than new inventions of machine learning. The contribution of this thesis is therefore not the discovery that ratios matter; it is the audited integration of ratio, fundamental, trend, macro, event, target, validation, and dashboard components.",
        "Dynamic and market-based default research supports repeated observations and forward horizons. Shumway (2001) criticizes static designs and argues that bankruptcy prediction should treat firms as observed repeatedly over time. Duffie, Saita and Wang (2007) use stochastic covariates in multi-period default prediction. Campbell, Hilscher and Szilagyi (2008) link distress risk to profitability, leverage, and market information. Merton (1974), Hillegeist et al. (2004), Bharath and Shumway (2008), and Bauer and Agarwal (2014) provide the structural, distance-to-default, and hazard-model comparison tradition. These works justify the firm-period perspective and future horizons, but they also mark a scope boundary: market-value and volatility features are not part of the final production model in this thesis.",
        "Macro-financial literature explains why market shifts belong in the design. Bernanke, Gertler and Gilchrist (1999) describe the financial accelerator, in which credit frictions amplify shocks. Gilchrist and Zakrajsek (2012) show that credit spreads contain business-cycle information, while Longstaff, Mithal and Neis (2005) caution that spread measures combine default and liquidity components. Tinoco and Wilson (2013) explicitly combine accounting, market, and macro variables in distress prediction. In this thesis, macro and global-event fields provide market-shift context and segmentation. They are not claimed to replace firm-level financial condition as the main empirical signal.",
        "Recent systematic reviews show that the field remains active and methodologically diverse. Billios, Seretidou and Stavropoulos (2024), Dasilas and Rigani (2024), Zhao, Ouenniche and De Smedt (2024), and Kim, Cho and Ryu (2020) document modern bankruptcy-prediction work using numerical indicators and machine learning. Bragoli et al. (2022) show that industry variables may matter. Tian and Yu (2017), Alaminos et al. (2016), Charalambakis and Garrett (2015), Hunter and Isachenkova (2001), Afanasev (2023), Fedorova, Musienko and Fedorov (2020), and Nevredinov (2021) also warn that country, sector, and institutional setting influence predictor behavior. These sources support the thesis boundary that U.S. SEC results cannot be mechanically generalized to private firms, Russia, emerging markets, or other accounting regimes.",
        "Machine-learning literature supports the model families and evaluation logic. Breiman (2001) provides the random forest foundation, Friedman (2001) the gradient boosting foundation, and Pedregosa et al. (2011) the scikit-learn implementation reference. Lessmann et al. (2015) and Barboza, Kimura and Altman (2017) justify comparing multiple model families rather than relying on one model. Rare-event evaluation is especially important because strict legal distress is sparse. Fawcett (2006) explains ROC analysis, while Davis and Goadrich (2006) and Saito and Rehmsmeier (2015) show why precision-recall metrics are more informative when positives are rare. Chawla et al. (2002) and Burez and Van den Poel (2009) show that class imbalance is a real modeling issue, although this thesis does not make synthetic oversampling the main empirical story.",
        "Interpretability research is necessary because predictive insights must be understandable. Lundberg and Lee (2017), Ribeiro, Singh and Guestrin (2016), Guidotti et al. (2018), Bussmann et al. (2021), Zhang et al. (2022), and Molnar (2025) show that model explanations can make complex predictors more usable. However, these works also imply a caution: explanation is not causality. The final empirical results in this thesis use bounded feature-family interpretation and do not claim that feature importance proves a causal mechanism. SHAP and LIME are reviewed as literature and future-work context, not as implemented final result methods.",
        "Design-science and business-analytics literature supports the dashboard artifact. Hevner et al. (2004) and Peffers et al. (2007) justify evaluating an information-system artifact, not only a statistical model. Chen, Chiang and Storey (2012) position business analytics as a field that links data, methods, and decision processes. Shneiderman (1996) and Arnott and Pervan (2014) support interactive decision-support design. The dashboard in this thesis is therefore a research artifact for historical exploration, model transparency, and caveat communication. It is not a governed production system or live legal-bankruptcy oracle.",
        "Missing-data literature is central because SEC accounting panels contain structural, reporting, and derivation gaps. Rubin (1976), Little and Rubin (2019), Graham (2009), and van Buuren (2018) show that missingness mechanisms affect inference and should not be ignored. This thesis preserves nulls, distinguishes unknown future labels from negatives, and excludes missingness indicators from substantive feature-group interpretation. Missingness indicators may help a model handle data structure, but they are not interpreted as economic causes of success or failure.",
    ]:
        add_para(doc, text)

    add_table(
        doc,
        "Table 2",
        "Literature synthesis and design response",
        ["Literature stream", "Representative authors", "Contribution to thesis", "Prior limitation", "Thesis response"],
        [
            ["Classical distress prediction", "Beaver; Altman; Ohlson; Zmijewski", "Supports accounting ratios, profitability, leverage, and liquidity.", "Often static and tied to one distress definition.", "Uses target hierarchy and firm-period panel."],
            ["Dynamic/default prediction", "Shumway; Duffie et al.; Campbell et al.", "Supports repeated observations and forward horizons.", "Often not connected to dashboard artifact or missing-aware targets.", "Uses filing-date timestamps and temporal validation."],
            ["Macro-finance and market shifts", "Bernanke et al.; Gilchrist and Zakrajsek; Tinoco and Wilson", "Supports macro regimes, spreads, rates, inflation, and financial conditions.", "Macro variables can be detached from firm-specific interpretation.", "Combines firm features with macro regimes and event windows."],
            ["Machine learning and rare events", "Breiman; Friedman; Davis and Goadrich; Saito and Rehmsmeier", "Supports tabular ML, PR-AUC, F1, ranking, and imbalanced evaluation.", "High ROC-AUC may mislead under rare positives.", "Reports PR-AUC, F1, positives, calibration, and ranking diagnostics."],
            ["Interpretability and design science", "Lundberg and Lee; Ribeiro et al.; Hevner et al.; Peffers et al.", "Supports interpretation and artifact evaluation.", "Explanations can be mistaken for causality.", "Uses bounded feature-family interpretation and dashboard caveats."],
        ],
        "The table turns the literature review into a design justification rather than a descriptive source list.",
        "Author's synthesis from the cited literature.",
    )

    add_heading(doc, "2.1 Critical Synthesis of the Research Gap", 2)
    for text in [
        "The literature establishes strong foundations but does not remove the need for the present empirical design. Classical distress models demonstrate the usefulness of accounting signals, yet many classical designs were built around smaller samples, static comparisons, or single failure definitions. Dynamic default models solve part of the time dimension, but they often emphasize event prediction rather than dashboard-supported interpretation. Machine-learning studies improve classification power, but their results can become difficult to defend when target construction, timestamp alignment, and leakage controls are underdocumented.",
        "The first gap is target ambiguity. A thesis on success and failure factors cannot rely only on formal bankruptcy because bankruptcy is legally specific and sparse. It also cannot treat any weak financial ratio as failure because that would erase the distinction between pressure, distress, and ordinary variation. The target hierarchy responds to this gap by assigning different empirical roles to strict legal distress, broader conservative pressure, success/resilience, and secondary dimensional outcomes.",
        "The second gap is validation realism. Financial panels are time ordered, and firm observations are serially related. Random row splitting can make results look better than they would in a forward-looking setting because the model can learn from later market regimes or adjacent observations. The thesis therefore follows a temporal validation protocol in which train, validation, and test observations are separated by prediction date. This design is closer to the analytical situation in which a user has past filings and wants to evaluate later outcomes.",
        "The third gap is interpretation discipline. Interpretability literature offers many methods, but explanation outputs become dangerous when they are read as causal proof. This thesis uses economic feature-family interpretation because it is easier to defend at the level of business factors: fundamentals, ratios, trends, macro context, and industry context. Individual feature rankings remain useful diagnostics, but the thesis emphasizes groups because they connect model behavior to economic categories without overstating precision.",
        "The fourth gap is artifact integration. A model table alone does not show whether data coverage, target denominators, missing labels, feature groups, and caveats are understandable to a user. A dashboard can make these elements inspectable, but only if the dashboard presents limitations as part of the artifact. This is why the thesis treats the dashboard as a design-science output and not merely as a decorative visualization layer.",
    ]:
        add_para(doc, text)

    add_heading(doc, "3. Research Methodology", 1)
    add_heading(doc, "3.1 Research Design and Data Architecture", 2)
    for text in [
        "The methodology follows an applied predictive-analytics and design-science design. The empirical system builds a panel of firm-period observations, constructs future outcome labels, trains baseline models under temporal validation, audits the data and feature set, interprets model behavior through economic feature families, and presents the results in a dashboard artifact. The unit of observation is one firm, one accounting period, and one prediction date. The prediction date is the SEC filing date, which represents the information availability date used for prediction.",
        "The production panel uses four main source families. The first is SEC Financial Statement Data Sets and EDGAR metadata, which provide structured accounting facts and filing dates (U.S. Securities and Exchange Commission, 2024, 2026). The second is FRED and related official macro-financial time series, which provide rates, spreads, inflation, labor, demand, oil, volatility, and financial-conditions context (Federal Reserve Bank of St. Louis, 2026; Federal Reserve Bank of Chicago, 2026). The third is a curated global-event calendar that maps observed event windows to market-shift context. The fourth is a manually maintained and source-verified formal distress-event file used for strict legal distress labels. These are joined with industry and firm metadata for sector filtering and dashboard interpretation.",
        "SEC accounting facts are processed conservatively. Balance-sheet variables use point-in-time values. Flow variables use current-period values when reported and avoid selecting prior-year comparative facts as current-period facts. The selected-fact provenance audit verifies current-period alignment and selected-value consistency. The final acceptance audit reports 700,844 selected SEC fact rows, zero selected rows with ddate not equal to period, and zero selected-value mismatches above tolerance. This does not mean that SEC taxonomy harmonization is perfect; it means that the selected production values are auditable under the implemented conservative mapping rules.",
        "The information-date logic is one of the most important controls in the empirical design. A fiscal period end date indicates the period covered by the report, but it is not the date when an analyst could have used the information. The production panel therefore defines prediction_date as the SEC filing date for every observation. This prevents the model from treating a December fiscal period end as if the full filed report were already observable on that date.",
        "The panel also distinguishes stable identifiers from display identifiers. CIK is treated as the stable SEC company key. Tickers are useful for readability, dashboard filters, and event-source matching, but they can change, delist, or function as aliases. The PCG/PGNPQ duplicate-CIK issue was resolved by preserving one CIK history and retaining alias metadata for event lookup. This matters because duplicate histories with different labels would create artificial evidence and contaminate both counts and model training.",
        "The macro layer is joined by prediction date and is interpreted as context available at or before the information date. It includes rate, spread, volatility, inflation, labor, demand, housing, credit, oil, dollar, and financial-conditions variables, as well as regime indicators derived from those series. The global-event layer is a curated historical context layer. It helps describe periods such as pandemic stress, banking stress, oil shocks, geopolitical war, monetary inflation, sovereign debt stress, and supply-chain disruption, but it is not a live event feed.",
    ]:
        add_para(doc, text)

    add_table(
        doc,
        "Table 3",
        "Data sources and methodological role",
        ["Source family", "Role in thesis", "Important limitation"],
        [
            ["SEC Financial Statement Data Sets", "Accounting fundamentals, filing dates, and structured firm-period facts.", "Requires conservative concept mapping, period selection, and provenance audits."],
            ["SEC EDGAR metadata", "Filing information date and identifier context.", "Filing availability differs from fiscal period end date."],
            ["FRED and macro-financial series", "Rates, spreads, inflation, labor, demand, credit, oil, VIX, and financial-conditions context.", "Contextual variables, not causal firm-level explanations by themselves."],
            ["Curated global-event calendar", "Historical market-shift windows for dashboard and sector/event comparison.", "Not a full live news or event database."],
            ["Distress-event date file", "Strict legal distress labels and forward event windows.", "Clean benchmark, but rare and not a universal bankruptcy database."],
            ["Industry and firm metadata", "Sector filters, CIK/ticker identifiers, SIC, and dashboard context.", "Identifiers and sampling metadata are excluded from economic prediction features when inappropriate."],
        ],
        "Data sources are used only within the empirical scope actually implemented.",
        "Author's synthesis from project data documentation and official source documentation.",
    )

    add_table(
        doc,
        "Table 4",
        "Current production panel",
        ["Item", "Current value"],
        [
            ["Rows", "31,702"],
            ["Columns", "190"],
            ["SEC CIKs", "540"],
            ["Ticker/display identifiers", "540"],
            ["Period-date range", "2009-03-31 to 2026-02-28"],
            ["Prediction-date range", "2009-04-15 to 2026-03-31"],
            ["Prediction-date policy", "prediction_date = filed_date for all rows"],
        ],
        "These are the accepted May 12, 2026 production values. Older 260-firm, 16,640-row, and random-split values are historical and are not used in the final empirical claims.",
        "THESIS_PROJECT_AUDIT_REVIEW_20260512.md and FINAL_SOURCE_OF_TRUTH_INDEX.md.",
    )

    add_heading(doc, "3.2 Target Construction and Missingness Logic", 2)
    for text in [
        "Target construction is the central methodological choice after timestamp alignment. The strict legal distress target uses formal distress event dates after the prediction date and within 456 days. Broader conservative failure pressure uses strict distress or repeated future financial weakness with balance stress, deterioration, or unhealthy-future signals. Success/resilience uses future observations with positive net income, positive ROA, acceptable leverage/assets, and no formal distress. Secondary outcomes are included to support dimensional analysis, not to replace the three primary targets.",
        "The suffixes in target names must be read carefully. The strict distress target uses a four-quarter legal-event horizon operationalized as 456 calendar days. The broader and secondary next_4obs targets evaluate the next four future firm observations. The success_resilience_next_4q name is retained from the production workflow, but the implemented missing-aware logic evaluates the next four future firm observations. Missing future information is preserved as unknown when a valid label cannot be assigned.",
        "The broader failure-pressure target is intentionally conservative. It does not label any temporary weak quarter as failure pressure. Positive labels require formal distress or repeated future profitability weakness combined with serious balance stress, deterioration, or unhealthy-future evidence. This makes the target broader than legal distress but narrower than loose financial weakness. The label is therefore suitable for failure-factor analysis, but it is not a bankruptcy label.",
        "The success/resilience target is also conservative in a different way. It requires enough observed future components to judge the label. If future net income, ROA, and leverage/assets are not sufficiently observed, the target remains missing rather than becoming an automatic non-success label. This missing-aware logic prevents missing leverage or accounting fields from being silently converted into negative success outcomes.",
        "Secondary targets are included because success and failure are multidimensional. Industry-relative resilience asks whether a firm performs well compared with sector-year peers. Stress resilience restricts attention to elevated market-stress observations. Recovery starts from weak current conditions. Quality-success-cashflow adds operating-cash-flow support to the success idea. These targets enrich the interpretation but remain secondary to the three primary outcomes.",
    ]:
        add_para(doc, text)

    add_table(
        doc,
        "Table 5",
        "Target hierarchy and label counts",
        ["Target", "Known rows", "Positives", "Missing rows", "Role"],
        [
            ["distress_next_4q", "31,702", "207", "0", "Strict legal distress benchmark"],
            ["failure_pressure_conservative_v2_next_4obs", "29,805", "3,003", "1,897", "Main broader failure-pressure target"],
            ["success_resilience_next_4q", "20,152", "12,519", "11,550", "Main success/resilience target"],
            ["industry_relative_resilience_next_4obs", "20,039", "8,093", "11,663", "Validated secondary production outcome"],
            ["stress_resilience_next_4obs", "5,901", "3,707", "25,801", "Validated secondary production outcome"],
            ["recovery_next_4obs", "11,375", "5,578", "20,327", "Validated secondary production outcome"],
            ["quality_success_cashflow_next_4obs", "16,767", "9,852", "14,935", "Validated secondary production outcome"],
        ],
        "Known rows exclude target-missing observations for each label. Missing labels are validity-preserving unknowns, not automatic negatives.",
        "Author's calculations from the production panel and final source-of-truth documentation.",
    )

    add_heading(doc, "3.3 Temporal Validation and Modeling Protocol", 2)
    add_table(
        doc,
        "Table 6",
        "Primary temporal-validation split",
        ["Split", "Prediction-date range", "Target", "Rows", "Positive cases", "Purpose"],
        [
            ["Train", "2009-04-15 to 2018-12-27", "distress_next_4q", "17,145", "58", "Model fitting"],
            ["Train", "2009-04-15 to 2018-12-27", "failure_pressure_conservative_v2_next_4obs", "17,156", "1,557", "Model fitting"],
            ["Train", "2009-04-15 to 2018-12-21", "success_resilience_next_4q", "11,064", "7,034", "Model fitting"],
            ["Validation", "2019-01-08 to 2021-12-22", "distress_next_4q", "6,045", "53", "Model and threshold selection"],
            ["Validation", "2019-01-08 to 2021-12-22", "failure_pressure_conservative_v2_next_4obs", "6,032", "836", "Model and threshold selection"],
            ["Validation", "2019-01-08 to 2021-12-22", "success_resilience_next_4q", "4,192", "2,415", "Model and threshold selection"],
            ["Test", "2022-01-05 to 2024-12-20", "distress_next_4q", "5,872", "96", "Final forward-time evaluation"],
            ["Test", "2022-01-05 to 2024-12-20", "failure_pressure_conservative_v2_next_4obs", "5,735", "575", "Final forward-time evaluation"],
            ["Test", "2022-01-05 to 2024-12-20", "success_resilience_next_4q", "4,047", "2,648", "Final forward-time evaluation"],
        ],
        "The split is based on prediction_date, not period_date. Later observations are not used to train earlier predictions. Counts are target-specific because missing-aware targets have different known-label denominators.",
        "Saved temporal model outputs in reports/modeling/panel_v2_*_model_metrics.csv and target-definition review notes.",
    )

    add_heading(doc, "3.4 Evaluation, Leakage Control, and Dashboard Design", 2)
    for text in [
        "This temporal split is deliberately different from a random 80/20 split. A random row split can allow the model to learn from later periods while being evaluated on earlier or adjacent observations from the same time environment. That design is especially risky in repeated firm-period panels because the same company can appear many times. The final thesis therefore evaluates forward-time generalization: earlier prediction dates train the model, the intermediate period selects thresholds and model choices, and the later period is held for final evaluation.",
        "The baseline model families are intentionally standard and reproducible: logistic regression as a linear reference, random forest for nonlinear tabular interactions, and gradient boosting for boosted-tree comparison. Numeric features are imputed within the training pipeline using training-learned medians; categorical features are encoded within the pipeline. Missingness indicators created by the imputer may support prediction, but they are excluded from substantive feature-family interpretation. The final production workflow does not rely on AutoGluon as the main empirical model.",
        "Evaluation metrics are matched to target prevalence. ROC-AUC is reported as a ranking measure across thresholds, but PR-AUC, precision, recall, F1, confusion matrices, calibration, and top-rank diagnostics are more informative for rare or imbalanced outcomes. The strict distress target is especially rare, so its PR-AUC and top-rank enrichment are more meaningful than any claim that the model is a general bankruptcy oracle.",
        "Leakage control is treated as part of methodology. The feature blacklist excludes target columns, future outcome columns, post-event indicators, event dates, days-to-event fields, future observations, sample-design labels, identifiers not intended as economic predictors, and dashboard-only explanatory fields. The May 12 audit reports that model feature-list leakage checks passed for the production targets. Macro and global-event timing audits also passed, supporting the claim that future market information is not used as a predictor.",
        "The dashboard artifact follows design-science logic. It exposes dataset composition, target coverage, data missingness, model metrics, feature groups, firm trajectories, macro and event comparisons, and caveat panels. Its purpose is transparent decision support and historical exploration. It is not a production risk engine, live SEC monitor, or legal-bankruptcy oracle.",
        "The model-output protocol also separates model selection from final evaluation. Candidate models and thresholds are selected on the validation period. Final test metrics are reported after those choices are fixed. Calibration diagnostics use validation-period predictions to choose calibration and then report test-period calibration outputs. This ordering protects the test period from becoming a tuning set.",
        "Feature interpretation is organized at two levels. The first level is economic feature-family importance, which is the primary interpretation used in the thesis. The second level is detailed feature and reason-code output for secondary diagnostics. The thesis emphasizes the first level because it is stable enough for academic interpretation and less likely to invite causal overreading than isolated single-feature ranks.",
    ]:
        add_para(doc, text)

    add_heading(doc, "4. Results", 1)
    add_heading(doc, "4.1 Panel Integrity and Target Coverage", 2)
    for text in [
        "The results are presented in the order required for empirical interpretation: panel integrity, target prevalence, baseline temporal model metrics, target-specific interpretation, calibration and ranking diagnostics, feature-family interpretation, and dashboard evidence. This sequence matters because model scores are only meaningful after the reader understands what data were built, what labels mean, and how the test period was protected from future information.",
        "The first result is structural. The accepted production panel contains 31,702 rows, 190 columns, 540 SEC CIKs, and 540 ticker/display identifiers. Prediction dates are filing dates for all rows. Duplicate ticker-period-prediction and CIK-period-prediction rows are zero after the duplicate-handling fixes. Post-event rows are retained for historical dashboard context but excluded from primary forward-looking model training. This supports the audit-readiness of the empirical system.",
        "The second result is target prevalence. Strict legal distress has 207 positives among 31,702 known rows, which makes it a rare-event benchmark. Broader conservative failure pressure has 3,003 positives among 29,805 known rows, making it better suited for failure-factor analysis. Success/resilience has 12,519 positives among 20,152 known rows, making it a positive financial-resilience counterpart with a very different denominator and class balance. The secondary targets add interpretable dimensions, but they remain secondary production outcomes.",
        "The denominator differences are substantive results rather than technical clutter. A strict legal distress target can be known for every row because a formal event either falls within the 456-day window or does not. A missing-aware success target requires future observations with enough financial components. A secondary stress-resilience target is only meaningful when current stress is present and future health can be observed. The thesis therefore reports known rows, missing rows, positives, and test positives rather than only percentages.",
        "The audit outputs also support the distinction between structural readiness and empirical validity. Passing file parsing, leakage, and timestamp audits does not prove the economic truth of the model. It shows that the empirical artifact is coherent enough for model evaluation and thesis interpretation. This distinction is important because data-engineering success is a necessary condition for the thesis, not the final research result by itself.",
    ]:
        add_para(doc, text)

    add_heading(doc, "4.2 Temporal Model Metrics", 2)
    add_table(
        doc,
        "Table 7",
        "Baseline temporal test metrics",
        ["Target", "Best baseline model", "Test rows", "Test positives", "Test PR-AUC", "Test F1"],
        [
            ["distress_next_4q", "Gradient boosting", "5,872", "96", "0.104", "0.182"],
            ["failure_pressure_conservative_v2_next_4obs", "Random forest", "5,735", "575", "0.758", "0.725"],
            ["success_resilience_next_4q", "Random forest", "4,047", "2,648", "0.975", "0.938"],
            ["industry_relative_resilience_next_4obs", "Gradient boosting", "4,098", "1,678", "0.855", "0.787"],
            ["stress_resilience_next_4obs", "Random forest", "3,028", "1,920", "0.974", "0.937"],
            ["recovery_next_4obs", "Random forest", "2,419", "1,199", "0.977", "0.915"],
            ["quality_success_cashflow_next_4obs", "Gradient boosting", "3,998", "2,458", "0.953", "0.920"],
        ],
        "Metrics are test-period values from the accepted baseline model outputs. Secondary targets are interpreted as dimensional outcomes, not replacements for the primary hierarchy.",
        "Author's calculations from reports/modeling/panel_v2_*_model_metrics.csv and best-model JSON files.",
    )

    for text in [
        "Strict legal distress should be interpreted as a rare-event benchmark rather than as the main headline prediction task. In the temporal test period, the accepted gradient-boosting baseline records PR-AUC of 0.104 and F1 of 0.182 with 96 positive cases. These values reflect the difficulty of predicting formal legal distress events in a broad public-company panel. The result is therefore not presented as a high-recall bankruptcy alarm. Its value is that it anchors the target hierarchy to a clean legal event definition and shows whether top-ranked observations are enriched relative to the rare-event base rate.",
        "Broader conservative failure pressure is the main failure-factor result. The random forest baseline reaches test PR-AUC of 0.758 and F1 of 0.725 with 575 positives in the test period. This result is materially stronger than strict distress because the target is more common and because it captures repeated serious financial pressure that does not necessarily culminate in formal legal distress. The correct interpretation is not that the target is bankruptcy. It is a conservative economic pressure label suitable for factor analysis.",
        "Success/resilience is the strongest primary target. The random forest baseline reaches test PR-AUC of 0.975 and F1 of 0.938 with 2,648 positives in the test period. This target is not rare, and its score should not be compared mechanically with strict distress. Its role is to provide a positive financial-resilience counterpart: future observations with positive net income, positive ROA, acceptable leverage/assets, and no formal distress under missing-aware label logic.",
        "The four validated secondary production outcomes add dimensional evidence. Industry-relative resilience captures peer-relative performance and leverage standing. Stress resilience focuses on observations in elevated market stress. Recovery asks whether weak current observations later move toward healthier future states. Quality-success-cashflow requires resilience supported by positive operating cash flow. Their strong test metrics support dashboard exploration and robustness of the target hierarchy, but they do not prove causal turnaround mechanisms.",
        "The results also show why ROC-AUC cannot be the only metric. For rare targets, a model can rank some positives higher than negatives while still producing limited precision or recall at a decision threshold. PR-AUC, F1, and top-rank enrichment make this tradeoff visible. For common positive targets, very high PR-AUC must be read together with the base rate and target definition. This metric discipline is central to the thesis because the three primary targets differ sharply in prevalence.",
        "The accepted baselines are deliberately not presented as exhaustive algorithm search. Random forest and gradient boosting are strong tabular baselines, and advanced boosted-tree robustness exists in the project outputs, but the thesis argument does not depend on claiming that every possible model family was tried. The empirical contribution is the audited target-and-validation framework and its interpretable results, not a leaderboard competition.",
    ]:
        add_para(doc, text)

    add_table(
        doc,
        "Table 8",
        "P0 audit summary",
        ["Audit", "Status", "Interpretation"],
        [
            ["Prediction timestamp", "PASS", "All rows use filing-date prediction timestamps in the current panel."],
            ["Distress event dates", "PASS", "Formal strict-distress rows have source-provenance records."],
            ["Leakage", "PASS", "Target, future, event, and metadata leakage fields are excluded from model features."],
            ["Macro lag", "PASS", "Macro variables are aligned to prediction dates."],
            ["Global-event timing", "PASS", "Event fields are attached through observed event windows."],
            ["Post-event rows", "PASS", "Post-event rows are retained for history but excluded from primary forward models."],
        ],
        "The audit results support structural validity and timing discipline; they do not by themselves prove causal validity.",
        "THESIS_PROJECT_AUDIT_REVIEW_20260512.md and reports/data_quality/p0_audit_summary.md.",
    )

    add_heading(doc, "4.3 Interpretation, Calibration, and Dashboard Result", 2)
    add_table(
        doc,
        "Table 9",
        "Economic feature-group interpretation for primary targets",
        ["Target", "Model", "Feature group", "Share excluding missingness"],
        [
            ["distress_next_4q", "Gradient boosting", "Accounting fundamentals", "53.2%"],
            ["distress_next_4q", "Gradient boosting", "Firm trend / deterioration", "27.9%"],
            ["distress_next_4q", "Gradient boosting", "Financial ratios", "16.7%"],
            ["failure_pressure_conservative_v2_next_4obs", "Random forest", "Accounting fundamentals", "36.9%"],
            ["failure_pressure_conservative_v2_next_4obs", "Random forest", "Firm trend / deterioration", "34.6%"],
            ["failure_pressure_conservative_v2_next_4obs", "Random forest", "Financial ratios", "27.3%"],
            ["success_resilience_next_4q", "Random forest", "Financial ratios", "43.2%"],
            ["success_resilience_next_4q", "Random forest", "Firm trend / deterioration", "29.1%"],
            ["success_resilience_next_4q", "Random forest", "Accounting fundamentals", "26.2%"],
        ],
        "Shares are aggregated by economic feature family and exclude missingness indicators from substantive interpretation.",
        "reports/modeling/panel_v2_feature_contribution_economic_groups.csv.",
    )

    add_table(
        doc,
        "Table 10",
        "Calibration and ranking interpretation",
        ["Target", "Calibration or ranking evidence", "Interpretation"],
        [
            ["distress_next_4q", "Calibrated test PR-AUC about 0.100; top 1% precision 18.6% and 11.4x lift.", "Rare formal events remain difficult, but top-ranked observations are enriched relative to the base rate."],
            ["failure_pressure_conservative_v2_next_4obs", "Calibrated test PR-AUC about 0.739; top 10% captures 73.7% of positives with 73.9% precision.", "Useful alert-list behavior for broader failure-pressure review."],
            ["success_resilience_next_4q", "Calibrated test PR-AUC about 0.972 and test Brier 0.0615.", "Strong positive financial-resilience counterpart under missing-aware labels."],
        ],
        "Calibrated probabilities are validation-calibrated model outputs, not true live event probabilities.",
        "CALIBRATION_AND_RANKING_NOTE.md and validation-extension gate outputs.",
    )

    for text in [
        "Feature-family interpretation answers the substantive factor question more defensibly than isolated feature ranking. Across the primary targets, accounting fundamentals, financial ratios, and firm trend/deterioration features dominate the economic interpretation layer. This is consistent with classical financial-distress literature and modern numerical-indicator reviews. Macro conditions, regimes, and global-event variables still have analytical value because they provide the market-shift context in which firm condition is observed, but the current results do not support a claim that macro variables replace firm fundamentals as the main predictive signal.",
        "The dashboard result is practical rather than purely statistical. The artifact presents dataset and target composition, data coverage and missingness, primary and secondary model metrics, feature groups, firm explorer views, macro/event comparisons, target timelines, and caveat panels. It makes the target hierarchy visible and helps users separate raw accounting values, target labels, model scores, feature groups, and claim boundaries. This supports decision-making transparency, but it does not provide live legal-bankruptcy prediction or causal proof.",
        "The empirical answer is therefore balanced. Strict legal distress is clean, rare, and difficult. Broader failure pressure is the main failure-factor outcome because it has enough positives and a conservative economic interpretation. Success/resilience is the strongest positive counterpart. The secondary targets extend the analysis into relative resilience, stress resilience, recovery, and cash-flow-supported quality. The main associated factor groups are financial and accounting in nature, with macro and event windows providing context and segmentation.",
        "Ablation and stability evidence is consistent with the feature-family interpretation. Firm-only and firm-plus-context specifications both provide useful signal, but the strongest economic interpretation remains linked to firm fundamentals, ratios, and deterioration. Macro and event features are still useful for contextualizing market shifts, especially inside the dashboard, but the thesis avoids claiming that they dominate the firm layer.",
        "The dashboard evidence closes the empirical loop because it makes the results inspectable. A reader can move from aggregate target counts to model metrics, from model metrics to feature groups, from feature groups to firm trajectories, and from firm trajectories to macro/event context. The artifact also keeps caveats visible, which is necessary because the same output can otherwise be mistaken for causal diagnosis or live operational risk scoring.",
    ]:
        add_para(doc, text)

    add_heading(doc, "5. Discussion", 1)
    add_heading(doc, "5.1 Implications, Limitations, and Future Research", 2)
    for text in [
        "The first implication is that the thesis should not be presented as a simple bankruptcy-prediction success story. The strict legal distress target is important precisely because it is clean, rare, and difficult. Its modest PR-AUC and F1 are not hidden; they are interpreted as evidence that formal legal distress is a hard benchmark in a broad public-company panel. This protects the thesis from overclaiming and makes the broader target hierarchy necessary rather than cosmetic.",
        "The second implication is that target hierarchy improves scientific honesty. A single binary success/failure label would mix formal legal events, repeated financial weakness, ordinary profitability, sector-relative performance, stress survival, and recovery into one ambiguous category. The hierarchy separates these meanings. As a result, performance differences across targets become interpretable: target definition changes prevalence, denominator, learnability, and managerial meaning.",
        "The third implication is that firm-level accounting and deterioration signals remain central. The feature-group outputs show that accounting fundamentals, ratios, and trends are the main economic interpretation groups for the primary tree-based models. This does not prove causality. It does show that the model behavior is economically plausible and not driven by hidden identifiers or target leakage. For managers and analysts, this means that basic financial discipline remains visible in predictive systems: profitability, leverage, liquidity, cash availability, balance-sheet structure, operating performance, and deterioration trends matter for model association.",
        "The fourth implication is that macro regimes and global events should be framed as context. A high-rate regime, oil shock, pandemic window, banking-stress episode, or supply-chain event does not affect every firm equally. The dashboard helps compare firms and sectors within these environments, but the current evidence does not justify a claim that market-shift variables dominate firm condition. Their main value is segmentation, historical comparison, and interpretation of the environment surrounding the firm-level signal.",
        "The managerial value of the artifact is bounded but concrete. A manager, analyst, investor, or supervisor can use the dashboard to inspect target coverage, compare strict distress with broader pressure and resilience, review model metrics, identify top feature groups, and examine firm histories under macro contexts. The artifact supports triage and explanation. It does not replace detailed due diligence, legal analysis, credit underwriting, or managerial judgment.",
        "For researchers, the most transferable contribution is the target-design discipline. The same hierarchy could be adapted to other jurisdictions, private-company datasets, banking portfolios, or sector-specific studies, but only if outcome definitions, timestamps, missingness rules, and leakage controls are rebuilt for the new setting. The U.S. SEC/FRED design is reproducible and useful, but it is not automatically portable to Russian firms, emerging markets, or private-company data.",
        "Several limitations remain. First, the panel is U.S. public-company oriented and reflects SEC reporting rules, public-firm disclosure incentives, and U.S. market institutions. Second, strict distress event dates are source-verified for the current formal rows, but the file is not a universal legal bankruptcy database. Third, SEC concept mapping is conservative and audited, but not perfect taxonomy harmonization. Fourth, missingness remains material for some accounting concepts and targets. Fifth, feature importance and feature-family shares are model-behavior evidence, not causal mechanisms. Sixth, the dashboard is a research artifact, not a governed production system.",
        "Future research can extend the framework in several directions. Market-price variables and distance-to-default measures would connect the accounting panel with structural credit-risk literature, provided that timestamp alignment and coverage are maintained. Textual disclosures, earnings-call language, ESG ratings, patents, HR measures, and customer sentiment could add richer explanatory dimensions, but none of these should be claimed without implementation and audit. Formal interpretability methods such as SHAP, permutation importance, partial dependence, and stability testing could deepen explanation in later versions. User studies could evaluate whether the dashboard improves analyst understanding of target definitions, uncertainty, and caveats.",
    ]:
        add_para(doc, text)

    add_heading(doc, "Conclusion", 1)
    for text in [
        "The thesis set out to determine how a reproducible, temporally valid, and interpretable predictive-analytics framework can identify and compare observable success and failure factors across industries under market-shift conditions. The answer is that such a framework must combine audited data construction, filing-date prediction timestamps, target hierarchy, temporal validation, leakage controls, missing-aware labels, feature-family interpretation, and transparent dashboard presentation. The contribution is not a claim of causal discovery or a universal bankruptcy model. It is a disciplined business-analytics artifact for predictive association and decision support.",
        "RQ1 asked which economic feature groups are most strongly associated with strict legal distress, broader conservative failure pressure, and success/resilience. The answer is that accounting fundamentals, financial ratios, and firm trend/deterioration features carry the main economic interpretation signal across the primary targets. Missingness indicators are excluded from substantive interpretation. Macro regimes and global-event variables provide context and segmentation, but the current results do not show that they replace firm-level financial condition as the main predictive signal.",
        "RQ2 asked how target definition affects predictive performance and interpretation under temporal validation. The answer is that target definition strongly affects both performance and meaning. Strict legal distress is rare and difficult, with test PR-AUC of 0.104 and F1 of 0.182 in the accepted baseline. Broader conservative failure pressure is more suitable for failure-factor analysis, with test PR-AUC of 0.758 and F1 of 0.725. Success/resilience is the strongest positive counterpart, with test PR-AUC of 0.975 and F1 of 0.938. These results cannot be ranked as if they were the same outcome; each target has a different denominator, prevalence, and interpretation.",
        "RQ3 asked how macro regimes and global-event windows contextualize firm outcomes across industries. The answer is that they provide market-shift context for sector comparison, dashboard filtering, and historical interpretation. Interest rates, credit spreads, inflation, volatility, oil prices, financial conditions, banking stress, pandemic windows, supply-chain stress, and geopolitical events help describe the environment in which firm outcomes occur. They do not, in the current results, displace fundamentals, ratios, and deterioration as the main factor-interpretation layer.",
        "RQ4 asked how an audited dashboard artifact can support transparent interpretation of predictive outputs without implying causal proof. The answer is that the dashboard supports exploration of dataset composition, target coverage, firm trajectories, model metrics, feature groups, macro context, event windows, and caveats. It improves interpretability by separating data, labels, predictions, explanations, and limitations. Its claim boundary is explicit: it is a historical decision-support and research artifact, not a live bankruptcy oracle and not causal evidence.",
        "The final contribution is therefore bounded and defensible. The thesis delivers an audited SEC/FRED/global-event firm-period panel, a target hierarchy aligned with the success/failure research question, forward-time validation, leakage and timing audits, feature-family interpretation, and a dashboard artifact. These elements together show how predictive analytics can support disciplined analysis of success and failure factors across industries while preserving scientific caution.",
    ]:
        add_para(doc, text)

    add_heading(doc, "References", 1)
    references = [
        "Afanasev, V. (2023). Default prediction model for emerging capital market service companies. Journal of Corporate Finance Research, 17(1), 64-77. https://doi.org/10.17323/j.jcfr.2073-0438.17.1.2023.64-77",
        "Alaminos, D., del Castillo, A., and Fernandez, M. A. (2016). A global model for bankruptcy prediction. PLOS ONE, 11(11), e0166693. https://doi.org/10.1371/journal.pone.0166693",
        "Altman, E. I. (1968). Financial ratios, discriminant analysis and the prediction of corporate bankruptcy. The Journal of Finance, 23(4), 589-609. https://doi.org/10.1111/j.1540-6261.1968.tb00843.x",
        "Altman, E. I., and Narayanan, P. (1997). An international survey of business failure classification models. Financial Markets, Institutions and Instruments, 6(2), 1-57. https://doi.org/10.1111/1468-0416.00010",
        "Arnott, D., and Pervan, G. (2014). A critical analysis of decision support systems research revisited: The rise of design science. Journal of Information Technology, 29, 269-293. https://doi.org/10.1057/jit.2014.16",
        "Barboza, F., Kimura, H., and Altman, E. (2017). Machine learning models and bankruptcy prediction. Expert Systems with Applications, 83, 405-417. https://doi.org/10.1016/j.eswa.2017.04.006",
        "Barney, J. (1991). Firm resources and sustained competitive advantage. Journal of Management, 17(1), 99-120. https://doi.org/10.1177/014920639101700108",
        "Bauer, J., and Agarwal, V. (2014). Are hazard models superior to traditional bankruptcy prediction approaches? A comprehensive test. Journal of Banking and Finance, 40, 432-442. https://doi.org/10.1016/j.jbankfin.2013.12.013",
        "Beaver, W. H. (1966). Financial ratios as predictors of failure. Journal of Accounting Research, 4, 71-111. https://doi.org/10.2307/2490171",
        "Bernanke, B. S., Gertler, M., and Gilchrist, S. (1999). The financial accelerator in a quantitative business cycle framework. Handbook of Macroeconomics, 1, 1341-1393. https://doi.org/10.1016/S1574-0048(99)10034-X",
        "Bharath, S. T., and Shumway, T. (2008). Forecasting default with the Merton distance to default model. The Review of Financial Studies, 21(3), 1339-1369. https://doi.org/10.1093/rfs/hhn044",
        "Billios, D., Seretidou, D., and Stavropoulos, A. (2024). The power of numerical indicators in predicting bankruptcy: A systematic review. Journal of Risk and Financial Management, 17(10), 433. https://doi.org/10.3390/jrfm17100433",
        "Bragoli, D., Ferretti, C., Ganugi, P., Marseguerra, G., Mezzogori, D., and Zammori, F. (2022). Machine-learning models for bankruptcy prediction: Do industrial variables matter? Spatial Economic Analysis, 17(2), 1-22. https://doi.org/10.1080/17421772.2021.1977377",
        "Breiman, L. (2001). Random forests. Machine Learning, 45, 5-32. https://doi.org/10.1023/A:1010933404324",
        "Burez, J., and Van den Poel, D. (2009). Handling class imbalance in customer churn prediction. Expert Systems with Applications, 36(6), 10365-10375. https://doi.org/10.1016/j.eswa.2009.05.027",
        "Bussmann, N., Giudici, P., Marinelli, D., and Papenbrock, J. (2021). Explainable machine learning in credit risk management. Computational Economics, 57, 203-216. https://doi.org/10.1007/s10614-020-10042-0",
        "Campbell, J. Y., Hilscher, J., and Szilagyi, J. (2008). In search of distress risk. The Journal of Finance, 63(6), 2899-2939. https://doi.org/10.1111/j.1540-6261.2008.01416.x",
        "Charalambakis, E., and Garrett, I. (2015). On the prediction of financial distress in developed and emerging markets. Review of Quantitative Finance and Accounting, 45, 349-368. https://doi.org/10.1007/s11156-014-0492-y",
        "Chawla, N. V., Bowyer, K. W., Hall, L. O., and Kegelmeyer, W. P. (2002). SMOTE: Synthetic minority over-sampling technique. Journal of Artificial Intelligence Research, 16, 321-357. https://doi.org/10.1613/jair.953",
        "Chen, H., Chiang, R. H. L., and Storey, V. C. (2012). Business intelligence and analytics: From big data to big impact. MIS Quarterly, 36(4), 1165-1188. https://doi.org/10.2307/41703503",
        "Dasilas, A., and Rigani, A. (2024). Machine learning techniques in bankruptcy prediction: A systematic literature review. Expert Systems with Applications, 255, 124761. https://doi.org/10.1016/j.eswa.2024.124761",
        "Davis, J., and Goadrich, M. (2006). The relationship between precision-recall and ROC curves. Proceedings of the 23rd International Conference on Machine Learning, 233-240. https://doi.org/10.1145/1143844.1143874",
        "Duffie, D., Saita, L., and Wang, K. (2007). Multi-period corporate default prediction with stochastic covariates. Journal of Financial Economics, 83(3), 635-665. https://doi.org/10.1016/j.jfineco.2005.10.011",
        "Duchek, S. (2020). Organizational resilience: A capability-based conceptualization. Business Research, 13, 215-246. https://doi.org/10.1007/s40685-019-0085-7",
        "Fawcett, T. (2006). An introduction to ROC analysis. Pattern Recognition Letters, 27(8), 861-874. https://doi.org/10.1016/j.patrec.2005.10.010",
        "Federal Reserve Bank of Chicago. (2026). National Financial Conditions Index. https://www.chicagofed.org/research/data/nfci/current-data",
        "Federal Reserve Bank of St. Louis. (2026). FRED API. https://fred.stlouisfed.org/docs/api/fred/",
        "Federal Reserve Board. (2025). Financial Stability Report. https://www.federalreserve.gov/publications/financial-stability-report.htm",
        "Fedorova, E. A., Musienko, S. O., and Fedorov, F. Y. (2020). Analysis of the external factors influence on the forecasting of bankruptcy of Russian companies. St Petersburg University Journal of Economic Studies, 36(1), 117-133. https://doi.org/10.21638/spbu05.2020.106",
        "Friedman, J. H. (2001). Greedy function approximation: A gradient boosting machine. The Annals of Statistics, 29(5), 1189-1232. https://doi.org/10.1214/aos/1013203451",
        "Gilchrist, S., and Zakrajsek, E. (2012). Credit spreads and business cycle fluctuations. American Economic Review, 102(4), 1692-1720. https://doi.org/10.1257/aer.102.4.1692",
        "Graham, J. W. (2009). Missing data analysis: Making it work in the real world. Annual Review of Psychology, 60, 549-576. https://doi.org/10.1146/annurev.psych.58.110405.085530",
        "Guidotti, R., Monreale, A., Ruggieri, S., Turini, F., Giannotti, F., and Pedreschi, D. (2018). A survey of methods for explaining black box models. ACM Computing Surveys, 51(5), 93. https://doi.org/10.1145/3236009",
        "Hannan, M. T., and Freeman, J. (1984). Structural inertia and organizational change. American Sociological Review, 49(2), 149-164. https://doi.org/10.2307/2095567",
        "Hepfer, M., and Lawrence, T. B. (2022). The heterogeneity of organizational resilience. Organization Theory, 3(1), 26317877221074701. https://doi.org/10.1177/26317877221074701",
        "Hevner, A. R., March, S. T., Park, J., and Ram, S. (2004). Design science in information systems research. MIS Quarterly, 28(1), 75-105. https://doi.org/10.2307/25148625",
        "Hillegeist, S. A., Keating, E. K., Cram, D. P., and Lundstedt, K. G. (2004). Assessing the probability of bankruptcy. Review of Accounting Studies, 9, 5-34. https://doi.org/10.1023/B:RAST.0000013627.90884.b7",
        "Hunter, J., and Isachenkova, N. (2001). Failure risk: A comparative study of UK and Russian firms. Journal of Policy Modeling, 23(5), 511-521. https://doi.org/10.1016/S0161-8938(01)00064-3",
        "Kim, M., Cho, S., and Ryu, D. (2020). Corporate default predictions using machine learning: Literature review. Sustainability, 12(16), 6325. https://doi.org/10.3390/su12166325",
        "Lessmann, S., Baesens, B., Seow, H. V., and Thomas, L. C. (2015). Benchmarking state-of-the-art classification algorithms for credit scoring. European Journal of Operational Research, 247(1), 124-136. https://doi.org/10.1016/j.ejor.2015.05.030",
        "Linnenluecke, M. K. (2017). Resilience in business and management research. International Journal of Management Reviews, 19(1), 4-30. https://doi.org/10.1111/ijmr.12076",
        "Little, R. J. A., and Rubin, D. B. (2019). Statistical Analysis with Missing Data. Wiley. https://doi.org/10.1002/9781119482260",
        "Longstaff, F. A., Mithal, S., and Neis, E. (2005). Corporate yield spreads: Default risk or liquidity? The Journal of Finance, 60(5), 2213-2253. https://doi.org/10.1111/j.1540-6261.2005.00797.x",
        "Lundberg, S. M., and Lee, S. I. (2017). A unified approach to interpreting model predictions. Advances in Neural Information Processing Systems 30. https://papers.neurips.cc/paper/7062-a-unified-approach-to-interpreting-model-predictions",
        "Merton, R. C. (1974). On the pricing of corporate debt: The risk structure of interest rates. The Journal of Finance, 29(2), 449-470. https://doi.org/10.1111/j.1540-6261.1974.tb03058.x",
        "Molnar, C. (2025). Interpretable Machine Learning. https://christophm.github.io/interpretable-ml-book/",
        "Nevredinov, A. R. (2021). The instrumental machine learning methods for corporate bankruptcy prediction. Finance and Credit, 27(9), 2118-2138. https://doi.org/10.24891/fc.27.9.2118",
        "Ohlson, J. A. (1980). Financial ratios and the probabilistic prediction of bankruptcy. Journal of Accounting Research, 18(1), 109-131. https://doi.org/10.2307/2490395",
        "Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., Blondel, M., Prettenhofer, P., Weiss, R., Dubourg, V., VanderPlas, J., Passos, A., Cournapeau, D., Brucher, M., Perrot, M., and Duchesnay, E. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825-2830. https://www.jmlr.org/papers/v12/pedregosa11a.html",
        "Peffers, K., Tuunanen, T., Rothenberger, M. A., and Chatterjee, S. (2007). A design science research methodology for information systems research. Journal of Management Information Systems, 24(3), 45-77. https://doi.org/10.2753/MIS0742-1222240302",
        "Platt, H. D., and Platt, M. B. (1994). Bankruptcy prediction with real variables. Journal of Business Finance and Accounting, 21(4), 491-510. https://doi.org/10.1111/j.1468-5957.1994.tb00332.x",
        "Ribeiro, M. T., Singh, S., and Guestrin, C. (2016). Why should I trust you? Explaining the predictions of any classifier. Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, 1135-1144. https://doi.org/10.1145/2939672.2939778",
        "Rubin, D. B. (1976). Inference and missing data. Biometrika, 63(3), 581-592. https://doi.org/10.1093/biomet/63.3.581",
        "Saito, T., and Rehmsmeier, M. (2015). The precision-recall plot is more informative than the ROC plot when evaluating binary classifiers on imbalanced datasets. PLOS ONE, 10(3), e0118432. https://doi.org/10.1371/journal.pone.0118432",
        "Shneiderman, B. (1996). The eyes have it: A task by data type taxonomy for information visualizations. Proceedings of the IEEE Symposium on Visual Languages, 336-343. https://doi.org/10.1109/VL.1996.545307",
        "Shumway, T. (2001). Forecasting bankruptcy more accurately: A simple hazard model. The Journal of Business, 74(1), 101-124. https://doi.org/10.1086/209665",
        "S&P Global Market Intelligence. (2025). U.S. corporate bankruptcy filings. https://www.spglobal.com/market-intelligence/en/news-insights/latest-news-headlines/us-bankruptcy-tracker",
        "Taffler, R. J. (1983). The assessment of company solvency and performance using a statistical model. Accounting and Business Research, 13(52), 295-308. https://doi.org/10.1080/00014788.1983.9729767",
        "Teece, D. J., Pisano, G., and Shuen, A. (1997). Dynamic capabilities and strategic management. Strategic Management Journal, 18(7), 509-533. https://doi.org/10.1002/(SICI)1097-0266(199708)18:7%3C509::AID-SMJ882%3E3.0.CO;2-Z",
        "Tian, S., and Yu, Y. (2017). Financial ratios and bankruptcy predictions: An international evidence. International Review of Economics and Finance, 51, 510-526. https://doi.org/10.1016/j.iref.2017.07.025",
        "Tinoco, M. H., and Wilson, N. (2013). Financial distress and bankruptcy prediction among listed companies using accounting, market and macroeconomic variables. International Review of Financial Analysis, 30, 394-419. https://doi.org/10.1016/j.irfa.2013.02.013",
        "U.S. Securities and Exchange Commission. (2024). EDGAR application programming interfaces. https://www.sec.gov/search-filings/edgar-application-programming-interfaces",
        "U.S. Securities and Exchange Commission. (2026). Financial Statement Data Sets. https://www.sec.gov/data-research/sec-markets-data/financial-statement-data-sets",
        "van Buuren, S. (2018). Flexible Imputation of Missing Data. CRC Press. https://doi.org/10.1201/9780429492259",
        "Zhang, Z., Wu, C., Qu, S., and Chen, X. (2022). An explainable artificial intelligence approach for financial distress prediction. Information Processing and Management, 59(4), 102988. https://doi.org/10.1016/j.ipm.2022.102988",
        "Zhao, J., Ouenniche, J., and De Smedt, J. (2024). Survey, classification and critical analysis of the literature on corporate bankruptcy and financial distress prediction. Machine Learning with Applications, 15, 100527. https://doi.org/10.1016/j.mlwa.2024.100527",
        "Zmijewski, M. E. (1984). Methodological issues related to the estimation of financial distress prediction models. Journal of Accounting Research, 22, 59-82. https://doi.org/10.2307/2490859",
    ]
    for ref in references:
        add_numbered(doc, ref)

    add_landscape_section(doc)
    add_heading(doc, "Appendices", 1)

    add_heading(doc, "Appendix A. Target Definitions and Missingness Rules", 2)
    add_table(
        doc,
        "Appendix Table A1",
        "Formal target-definition table",
        ["Target", "Positive condition", "Negative condition", "Missing condition", "Forward horizon", "Interpretation", "Claim boundary"],
        [
            ["distress_next_4q", "Formal distress event after prediction_date and within 456 days.", "No formal distress event within the 456-day forward window.", "None in full panel; post-event rows are excluded from primary models.", "456 calendar days.", "Strict legal distress benchmark.", "Clean but rare; not a complete bankruptcy database."],
            ["failure_pressure_conservative_v2_next_4obs", "Formal distress or at least three future nonpositive net-income observations plus balance stress, deterioration, or unhealthy-future signal.", "Known future horizon and positive rule not met.", "Post-event rows or fewer than three future observations unless formal distress occurs.", "Next four future firm observations.", "Main broader failure-pressure target.", "Financial pressure, not legal bankruptcy."],
            ["success_resilience_next_4q", "At least three future healthy observations with positive net income, positive ROA, leverage/assets from 0 to 0.85, and no formal distress.", "At least three valid future health observations and positive rule not met.", "Fewer than three valid future health observations.", "Next four future firm observations.", "Main financial success/resilience counterpart.", "Financial resilience, not total organizational success."],
            ["industry_relative_resilience_next_4obs", "At least three future observations above sector-year profitability median and not in the worst leverage quartile.", "Known peer-relative future horizon and positive rule not met.", "Insufficient relative-profitability or leverage comparison data, or post-event rows.", "Next four future firm observations.", "Peer-relative resilience.", "Secondary dimensional outcome."],
            ["stress_resilience_next_4obs", "Current elevated stress and at least three future healthy observations without formal distress.", "Current elevated stress, known future horizon, and positive rule not met.", "No current stress condition, insufficient future health data, or post-event rows.", "Next four future firm observations.", "Resilience conditional on stress.", "Secondary dimensional outcome."],
            ["recovery_next_4obs", "Current weak condition followed by at least two future healthy, positive-income, positive-ROA, and leverage-acceptable observations without formal distress.", "Current weak condition, known future horizon, and positive rule not met.", "No current weakness, insufficient future health data, or post-event rows.", "Next four future firm observations.", "Recovery from weak current condition.", "Secondary dimensional outcome."],
            ["quality_success_cashflow_next_4obs", "At least three future observations that are healthy and have positive operating cash flow, without formal distress.", "Known future quality/cash-flow horizon and positive rule not met.", "Insufficient future net income, ROA, leverage, or operating-cash-flow observations, or post-event rows.", "Next four future firm observations.", "Cash-flow-supported financial success.", "Secondary dimensional outcome."],
        ],
        "The table formalizes the existing production target logic. Null labels preserve validity when future information is insufficient.",
        "scripts/sec_fsd/build_panel_v2.py, scripts/modeling/run_exploratory_target_lab.py, and TARGET_DEFINITIONS_FINAL_REVIEW.md.",
        size=APPENDIX_TABLE_SIZE,
    )

    add_heading(doc, "Appendix B. Feature-Family Dictionary", 2)
    add_table(
        doc,
        "Appendix Table B1",
        "Feature-family dictionary",
        ["Feature family", "Examples", "Interpretation rule"],
        [
            ["Accounting fundamentals", "total_assets; total_liabilities; current_assets; cash_equivalents; revenue; operating_income; net_income; operating_cash_flow; capital_expenditure", "Direct SEC-derived firm condition and scale variables."],
            ["Financial ratios", "ROA; leverage/assets; current_ratio; cash/assets; net_margin; operating_margin; gross_margin; R&D intensity", "Comparable firm condition relative to size or revenue."],
            ["Firm trend and deterioration", "Lags; four-observation changes; growth rates; prior negative-income counts", "Direction and persistence of firm-level change."],
            ["Macro and regime context", "Interest rates; Treasury spreads; VIX; credit spreads; NFCI; inflation; oil prices; yield-curve inversion; high-rate regime", "Market-shift context and segmentation, not standalone causality."],
            ["Global-event context", "Pandemic windows; banking stress; supply-chain stress; oil shock; geopolitical war; monetary inflation", "Discrete historical event windows for dashboard comparison."],
            ["Industry and metadata", "Sector; SIC; CIK; ticker/display identifier", "Sector context and identifiers; identifiers are not interpreted as economic predictors."],
            ["Excluded or auxiliary fields", "Targets; future outcomes; post-event indicators; source labels; missingness indicators", "Excluded from economic interpretation or from predictive features when leakage-prone."],
        ],
        "Predictive feature families are separated from identifiers, sampling metadata, and target-derived fields.",
        "Model feature lists, feature blacklist, and MODEL_INTERPRETATION_STABILITY_NOTE.md.",
        size=APPENDIX_TABLE_SIZE,
    )

    add_heading(doc, "Appendix C. Leakage and Audit Controls", 2)
    add_table(
        doc,
        "Appendix Table C1",
        "Leakage blacklist categories",
        ["Blacklisted category", "Reason for exclusion"],
        [
            ["Target columns", "Would directly reveal the outcome being predicted."],
            ["Future outcome columns", "Contain information from after the prediction date."],
            ["Post-event indicators", "Reveal whether a formal event has already occurred."],
            ["Distress event dates", "Encode future legal-event timing."],
            ["Days-to-event fields", "Directly reveal proximity to the target event."],
            ["Future observations", "Break forward-time prediction logic."],
            ["Sampling metadata", "Can reveal case-control construction rather than economics."],
            ["Source labels", "May expose how firms entered the sample."],
            ["Identifiers not intended as predictors", "Useful for joins and dashboard filters, not economic prediction."],
            ["Dashboard-only explanatory fields", "Used for presentation and caveats, not model training."],
        ],
        "The leakage audit checks model feature lists against the blacklist and passed for the production outputs.",
        "config/model_feature_blacklist.csv and THESIS_PROJECT_AUDIT_REVIEW_20260512.md.",
        size=APPENDIX_TABLE_SIZE,
    )

    add_heading(doc, "Appendix D. Temporal Validation Split", 2)
    add_table(
        doc,
        "Appendix Table D1",
        "Primary target temporal split detail",
        ["Split", "Prediction years", "Target", "Rows", "Positive cases", "Purpose"],
        [
            ["Train", "2009-2018", "distress_next_4q", "17,145", "58", "Fit model parameters."],
            ["Train", "2009-2018", "failure_pressure_conservative_v2_next_4obs", "17,156", "1,557", "Fit model parameters."],
            ["Train", "2009-2018", "success_resilience_next_4q", "11,064", "7,034", "Fit model parameters."],
            ["Validation", "2019-2021", "distress_next_4q", "6,045", "53", "Choose model and threshold."],
            ["Validation", "2019-2021", "failure_pressure_conservative_v2_next_4obs", "6,032", "836", "Choose model and threshold."],
            ["Validation", "2019-2021", "success_resilience_next_4q", "4,192", "2,415", "Choose model and threshold."],
            ["Test", "2022-2024", "distress_next_4q", "5,872", "96", "Final forward-time evaluation."],
            ["Test", "2022-2024", "failure_pressure_conservative_v2_next_4obs", "5,735", "575", "Final forward-time evaluation."],
            ["Test", "2022-2024", "success_resilience_next_4q", "4,047", "2,648", "Final forward-time evaluation."],
        ],
        "The split is by prediction_date year after post-event and known-label filters; it is not a random split.",
        "Saved model-output split rows and positives.",
        size=APPENDIX_TABLE_SIZE,
    )

    add_heading(doc, "Appendix E. Dashboard Artifact", 2)
    figure_dir = ROOT / "reports" / "figures" / "dashboard"
    figures = [
        ("Figure E1", "Dashboard overview and target composition", "overview.png", "Shows panel and target composition for the dashboard's selected model window; it helps compare strict distress, broader pressure, success/resilience, and unknown horizons, but it does not prove causal relationships."),
        ("Figure E2", "Data coverage and missingness", "data_coverage.png", "Shows coverage and missingness by feature family; it helps explain why nulls are preserved, but it does not imply that missingness is an economic cause."),
        ("Figure E3", "Model metrics view", "models_failure_pressure.png", "Shows model metrics for a primary target; it helps compare temporal model behavior, but it does not provide a live probability guarantee."),
        ("Figure E4", "Feature-group interpretation view", "target_lab_factor_groups.png", "Shows factor groups used for interpretation; it supports economic reading of model behavior, but it does not prove causality."),
        ("Figure E5", "Firm explorer view", "firm_explorer.png", "Shows firm-level trajectory and context; it supports case review, but it is not a substitute for due diligence."),
        ("Figure E6", "Macro/event comparison view", "macro_compare_training_baseline.png", "Shows macro comparison using a training-period baseline; it supports market-shift context, not macro-dominance claims."),
        ("Figure E7", "Artifact notes and caveats", "artifact_notes.png", "Shows scope and caveat language; it helps preserve the thesis claim boundary and prevents dashboard overinterpretation."),
    ]
    for number, title, filename, caption in figures:
        add_figure(doc, number, title, figure_dir / filename, caption)

    add_heading(doc, "Appendix F. Model-Output Support", 2)
    add_table(
        doc,
        "Appendix Table F1",
        "Accepted primary model-output support",
        ["Target", "Best model", "ROC-AUC", "PR-AUC", "Precision", "Recall", "F1"],
        [
            ["distress_next_4q", "Gradient boosting", "0.816", "0.104", "0.120", "0.375", "0.182"],
            ["failure_pressure_conservative_v2_next_4obs", "Random forest", "0.954", "0.758", "0.669", "0.791", "0.725"],
            ["success_resilience_next_4q", "Random forest", "0.967", "0.975", "0.941", "0.934", "0.938"],
        ],
        "Primary model outputs are selected from saved model files, with validation used for threshold/model selection and test used only for final evaluation.",
        "reports/modeling/panel_v2_*_best_model.json and model metrics CSV files.",
        size=APPENDIX_TABLE_SIZE,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
